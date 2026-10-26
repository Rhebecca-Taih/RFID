# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# BIBLIOTECAS
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
import os
import re
import sys
import json
import time
import shutil
import tempfile
import subprocess
import threading
import importlib
import webbrowser
from datetime import datetime, time as dtime, date as ddate

import tkinter as tk

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

pasta_script = os.path.dirname(os.path.abspath(__file__))
pasta_raiz   = os.path.dirname(pasta_script)
_NO_WINDOW   = getattr(subprocess, "CREATE_NO_WINDOW", 0)


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# DETECÇÃO PARA AUTO-INSTALAÇÃO DE BIBLIOTECAS
# Instalação ao vivo em _checar_bibliotecas (tela Bibliotecas).
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_LIBS = [
    ("openpyxl",      "openpyxl"),
    ("dotenv",        "python-dotenv"),
    ("selenium",      "selenium"),
    ("pystray",       "pystray"),
    ("PIL",           "Pillow"),
    ("customtkinter", "customtkinter"),
]
_STATUS_INSTALL = {}
for _ni, _np in _LIBS:
    try:
        __import__(_ni)
        _STATUS_INSTALL[_np] = "ok"
    except ImportError:
        _STATUS_INSTALL[_np] = "faltando"


if _STATUS_INSTALL.get("customtkinter") == "faltando":
    _splash = tk.Tk()
    _splash.title("Robô Aprovador")
    _splash.configure(bg="#18181B")
    _splash.resizable(False, False)
    _sw, _sh = _splash.winfo_screenwidth(), _splash.winfo_screenheight()
    _splash.geometry(f"360x110+{(_sw-360)//2}+{(_sh-110)//2}")
    tk.Label(_splash, text="Preparando o ambiente pela primeira vez...",
             font=("Segoe UI", 11), fg="#F2F2F2", bg="#18181B").pack(pady=(28, 6))
    tk.Label(_splash, text="Instalando customtkinter...",
             font=("Segoe UI", 9), fg="#A8A8A8", bg="#18181B").pack()
    _splash.update()
    try:
        subprocess.run([sys.executable, "-m", "pip", "install", "customtkinter", "--quiet"],
                       capture_output=True, timeout=180, creationflags=_NO_WINDOW)
        __import__("customtkinter")
        _STATUS_INSTALL["customtkinter"] = "ok"
    except Exception:
        _STATUS_INSTALL["customtkinter"] = "faltando"
    _splash.destroy()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# IMPORTS EXTERNOS
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_IMPORTS_OK = True
_IMPORTS_ERRO = ""

def _tentar_importar_reais():
    global _IMPORTS_OK, _IMPORTS_ERRO
    global ctk, load_dotenv, openpyxl, load_workbook, Workbook
    global webdriver, Options, Service, By, WebDriverWait, EC, Keys, TimeoutException, WebDriverException
    try:
        import customtkinter as ctk
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")
    except Exception as _ex:
        _IMPORTS_OK = False
        _IMPORTS_ERRO = f"customtkinter -- {_ex}"
        return

    try:
        from dotenv import load_dotenv
        import openpyxl
        from openpyxl import load_workbook, Workbook
        from selenium import webdriver
        from selenium.webdriver.chrome.options import Options
        from selenium.webdriver.chrome.service import Service
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
        from selenium.webdriver.common.keys import Keys
        from selenium.common.exceptions import TimeoutException, WebDriverException
        _IMPORTS_OK = True
        _IMPORTS_ERRO = ""
    except Exception as _ex:
        _IMPORTS_OK = False
        _IMPORTS_ERRO = str(_ex)

_tentar_importar_reais()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# CONFIGURAÇÕES
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
CAMINHO_ENV = os.path.join(pasta_raiz, "Config_Cg", ".env")

def _carregar_env():
    try:
        from dotenv import load_dotenv as _load_dotenv
        _load_dotenv(dotenv_path=CAMINHO_ENV)
        return
    except ImportError:
        pass
    try:
        with open(CAMINHO_ENV, "r", encoding="utf-8-sig") as f:
            for linha in f:
                linha = linha.strip()
                if not linha or linha.startswith("#") or "=" not in linha:
                    continue
                chave, valor = linha.split("=", 1)
                valor = valor.strip()
                if len(valor) >= 2 and valor[0] == valor[-1] and valor[0] in "\"'":
                    valor = valor[1:-1]
                os.environ.setdefault(chave.strip(), valor)
    except Exception:
        pass

_carregar_env()

PORTAL_USER      = os.getenv("PORTAL_USER", "")
PORTAL_SENHA     = os.getenv("PORTAL_SENHA", "")
PORTAL_URL_LOGIN = os.getenv("PORTAL_URL_LOGIN", "")
PORTAL_SITE      = os.getenv("PORTAL_SITE", "")
CAMINHO_REDE_BIN = os.getenv("CAMINHO_REDE_BIN", "")

CAMINHO_FILA      = os.path.join(pasta_script, "Fila_Aprovacao", "Fila_Aprovacao.xlsx")
CAMINHO_RESULTADO = os.path.join(pasta_script, "Fila_Aprovacao", "Resultado_Aprovacao.xlsx")


CAMINHO_STATUS_APROVADOR = os.path.join(pasta_script, "Fila_Aprovacao", "status_aprovador.txt")
PASTA_HISTORICO   = os.path.join(pasta_script, "Historico_Aprovacoes")


PASTA_HISTORICO_IMPRESSAO = os.path.join(pasta_raiz, "Impressão", "HISTORICO DE APROVAÇÃO")

ABA_ENTRADA = "APROVAÇÃO"
EXCEL_EPOCH = ddate(1899, 12, 30)


_LOCAL_BASE = os.path.join(os.environ.get("LOCALAPPDATA", pasta_script), "Aprovador_Bot")
PASTA_BIN   = os.path.join(_LOCAL_BASE, "Bin")
CHROME_PATH = os.path.join(PASTA_BIN, "chrome-win64",       "chrome.exe")
DRIVER_PATH = os.path.join(PASTA_BIN, "chromedriver-win64", "chromedriver.exe")


# XPaths em JSON externo > seletor ajustável quando o portal muda, sem mexer no código.
CAMINHO_XPATHS = os.path.join(pasta_script, "xpaths_aprovador.json")
try:
    with open(CAMINHO_XPATHS, "r", encoding="utf-8") as _f:
        XPATHS = json.load(_f)
except Exception:
    XPATHS = {}

def _recarregar_xpaths():
    global XPATHS
    try:
        with open(CAMINHO_XPATHS, "r", encoding="utf-8") as f:
            XPATHS = json.load(f)
    except Exception as e:
        registrar_erro_permanente(f"RECARREGAR XPATHS -- {e}")


# Crédito > aparece no canto superior direito do cabeçalho.
NOME_CRIADOR        = os.getenv("NOME_CRIADOR",        "Rhebecca-Taih")
LINK_GITHUB_CRIADOR = os.getenv("LINK_GITHUB_CRIADOR", "https://github.com/Rhebecca-Taih")


# Pausa noturna: encerra às 20h e retoma às 7h30, sem fechar o processo.
HORA_ENCERRAR = dtime(20, 0)
HORA_RETOMAR  = dtime(7, 30)


# Heartbeat a cada 10 min; lock expira após 2h sem atualização > PC travado não prende o lock.
INTERVALO_HEARTBEAT = 10 * 60
EXPIRACAO_LOCK       = 2 * 60 * 60
LIMITE_RETRY_NAVEGADOR_MORTO = 2

_status_atual = "Aguardando..."
_parar_loop   = threading.Event()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# PALETA VISUAL
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
COR_FUNDO        = "#18181B"
COR_CARD         = "#242428"
COR_HEADER       = "#FFC72C"
COR_HEADER_TXT   = "#1A1A1A"
COR_TEXTO        = "#F2F2F2"
COR_DETALHE      = "#A8A8A8"
COR_MUTED        = "#7A7A7A"
COR_BORDA        = "#3A3A3C"
COR_BTN_C        = "#333336"
COR_BTN_C_TXT    = "#F2F2F2"
COR_BTN_B        = "#3D8BFF"
COR_BTN_B_TXT    = "#FFFFFF"
COR_OK           = "#4ADE80"
COR_OK_FUNDO     = "#173321"
COR_AVISO        = "#FBBF24"
COR_AVISO_FUNDO  = "#3A2E0C"
COR_ERRO         = "#F87171"
COR_ERRO_FUNDO   = "#3A1414"


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# ARQUIVOS DE CONTROLE
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_pasta_logs    = os.path.join(pasta_script, "logs")
_pasta_runtime = os.path.join(pasta_script, "runtime")
try:
    os.makedirs(_pasta_logs,    exist_ok=True)
    os.makedirs(_pasta_runtime, exist_ok=True)
except Exception:
    pass

LOG_ARQUIVO          = os.path.join(_pasta_logs, "historico_aprovador.txt")
LOG_ERROS_PERMANENTE = os.path.join(_pasta_logs, "historico_erros.txt")
LOCK_ARQUIVO         = os.path.join(_pasta_runtime, "aprovador.lock")


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LOG
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _data_hoje():
    return datetime.now().strftime("%d/%m/%Y")

def _agora_completo():
    return datetime.now().strftime("%d/%m/%Y %H:%M:%S")

def _gerenciar_log():
    sep = "=" * 60
    hoje = _data_hoje()
    if os.path.exists(LOG_ARQUIVO):
        try:
            with open(LOG_ARQUIVO, "r", encoding="utf-8") as f:
                f.readline()
                segunda = f.readline()
            if hoje not in segunda:
                os.remove(LOG_ARQUIVO)
        except Exception:
            try:
                os.remove(LOG_ARQUIVO)
            except Exception:
                pass
    if not os.path.exists(LOG_ARQUIVO):
        try:
            with open(LOG_ARQUIVO, "w", encoding="utf-8") as f:
                f.write(f"{sep}\n  LOG DE AUTOMACAO -- {hoje}\n{sep}\n\n")
        except Exception:
            pass

PASTA_LOG_LOCAL = os.path.join(os.environ.get("LOCALAPPDATA") or os.path.expanduser("~"), "RobosBureau", "Aprovador")
DIAS_MANTER_LOG_LOCAL = 30
_rede_indisponivel_desde = None

def _limpar_logs_locais_antigos():
    try:
        limite = time.time() - DIAS_MANTER_LOG_LOCAL * 86400
        for nome in os.listdir(PASTA_LOG_LOCAL):
            caminho = os.path.join(PASTA_LOG_LOCAL, nome)
            if os.path.isfile(caminho) and os.path.getmtime(caminho) < limite:
                os.remove(caminho)
    except Exception:
        pass

def _registrar_log_local(texto):
    global _rede_indisponivel_desde
    if _rede_indisponivel_desde is None:
        _rede_indisponivel_desde = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    try:
        os.makedirs(PASTA_LOG_LOCAL, exist_ok=True)
        caminho = os.path.join(PASTA_LOG_LOCAL, f"log_{datetime.now().strftime('%d-%m-%Y')}.txt")
        arquivo_novo = not os.path.exists(caminho)
        with open(caminho, "a", encoding="utf-8") as f:
            f.write(texto)
        if arquivo_novo:
            _limpar_logs_locais_antigos()
    except Exception:
        pass

def _aviso_rede_restabelecida():
    if _rede_indisponivel_desde is None:
        return ""
    return (f"⚠️  Rede indisponível de {_rede_indisponivel_desde} até "
            f"{datetime.now().strftime('%d/%m/%Y %H:%M:%S')}. Registros desse período em: {PASTA_LOG_LOCAL}")

def _confirmar_rede_ok():
    global _rede_indisponivel_desde
    _rede_indisponivel_desde = None

def registrar_log(linhas):
    sep = "-" * 60
    texto_linhas = linhas if isinstance(linhas, list) else [str(linhas)]
    bloco = sep + "\n" + "".join(l + "\n" for l in texto_linhas) + "\n"
    try:
        _gerenciar_log()
        with open(LOG_ARQUIVO, "a", encoding="utf-8") as f:
            aviso = _aviso_rede_restabelecida()
            if aviso:
                f.write(f"{sep}\n{aviso}\n\n")
            f.write(bloco)
        _confirmar_rede_ok()
    except Exception:
        _registrar_log_local(f"[{_agora_completo()}]\n{bloco}")

def registrar_erro_permanente(msg):
    try:
        with open(LOG_ERROS_PERMANENTE, "a", encoding="utf-8") as f:
            f.write(f"[{_agora_completo()}] {msg}\n")
    except Exception:
        _registrar_log_local(f"[{_agora_completo()}] ERRO: {msg}\n")

def _gravar_status_ciclo(status, motivo=""):
    try:
        os.makedirs(os.path.dirname(CAMINHO_STATUS_APROVADOR), exist_ok=True)
        with open(CAMINHO_STATUS_APROVADOR, "w", encoding="utf-8") as f:
            f.write(f"{status}|{motivo}")
    except Exception:
        pass

def _ultimo_status_log():
    if not os.path.exists(LOG_ARQUIVO):
        return "Sem registros anteriores."
    try:
        with open(LOG_ARQUIVO, "r", encoding="utf-8") as f:
            linhas = [l.strip() for l in f.readlines()
                      if l.strip() and not l.strip().startswith("-")
                      and not l.strip().startswith("=")
                      and "LOG DE AUTOMACAO" not in l]
        return linhas[-1] if linhas else "Sem registros."
    except Exception:
        return "Erro ao ler o log."


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LOCK
# Formato do arquivo: PID|NOME_PC|DATA_INICIO|ULTIMO_HEARTBEAT
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _verificar_instancia_unica():
    nome_pc = os.environ.get("COMPUTERNAME", "DESCONHECIDO")
    if os.path.exists(LOCK_ARQUIVO):
        try:
            with open(LOCK_ARQUIVO, "r") as f:
                partes = f.read().strip().split("|")
            pid_antigo   = int(partes[0])
            pc_antigo    = partes[1] if len(partes) > 1 else "?"
            dt_inicio    = partes[2] if len(partes) > 2 else "?"
            dt_heartbeat = partes[3] if len(partes) > 3 else dt_inicio

            if pc_antigo == nome_pc:
                r = subprocess.run(f'tasklist /FI "PID eq {pid_antigo}" /NH',
                                   shell=True, capture_output=True, text=True,
                                   creationflags=_NO_WINDOW)
                if str(pid_antigo) in r.stdout:
                    return False, f"{pc_antigo} desde {dt_inicio}"
            else:
                try:
                    hb_dt = datetime.strptime(dt_heartbeat, "%d/%m/%Y %H:%M")
                    idade_seg = (datetime.now() - hb_dt).total_seconds()
                except Exception:
                    idade_seg = 0
                if idade_seg < EXPIRACAO_LOCK:
                    minutos_atras = max(0, int(idade_seg // 60))
                    return False, f"{pc_antigo} desde {dt_inicio} (última atividade há {minutos_atras} min)"
        except Exception:
            pass
        try:
            os.remove(LOCK_ARQUIVO)
        except Exception:
            pass
    try:
        agora_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        with open(LOCK_ARQUIVO, "w") as f:
            f.write(f"{os.getpid()}|{nome_pc}|{agora_str}|{agora_str}")
    except Exception:
        pass
    return True, ""

def _atualizar_heartbeat_lock():
    try:
        if not os.path.exists(LOCK_ARQUIVO):
            return
        with open(LOCK_ARQUIVO, "r") as f:
            partes = f.read().strip().split("|")
        if not partes or int(partes[0]) != os.getpid():
            return
        pc     = partes[1] if len(partes) > 1 else os.environ.get("COMPUTERNAME", "DESCONHECIDO")
        inicio = partes[2] if len(partes) > 2 else datetime.now().strftime("%d/%m/%Y %H:%M")
        with open(LOCK_ARQUIVO, "w") as f:
            f.write(f"{os.getpid()}|{pc}|{inicio}|{datetime.now().strftime('%d/%m/%Y %H:%M')}")
    except Exception:
        pass

def _liberar_lock():
    try:
        if os.path.exists(LOCK_ARQUIVO):
            with open(LOCK_ARQUIVO, "r") as f:
                pid = int(f.read().strip().split("|")[0])
            if pid == os.getpid():
                os.remove(LOCK_ARQUIVO)
    except Exception:
        pass


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# CHROME E BIN LOCAL
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _bin_completa():
    return os.path.exists(CHROME_PATH) and os.path.exists(DRIVER_PATH)

def _copiar_bin_da_rede():
    if _bin_completa():
        return True, ""
    if not CAMINHO_REDE_BIN:
        return False, "CAMINHO_REDE_BIN não configurado no .env"
    if not os.path.exists(CAMINHO_REDE_BIN):
        return False, f"Caminho de rede não encontrado: {CAMINHO_REDE_BIN}"
    try:
        os.makedirs(_LOCAL_BASE, exist_ok=True)
        shutil.copytree(CAMINHO_REDE_BIN, _LOCAL_BASE, dirs_exist_ok=True)
        if _bin_completa():
            return True, ""
        return False, "Cópia terminou mas os arquivos esperados não foram encontrados."
    except Exception as e:
        return False, str(e)

def _checar_chrome():
    r = {"chrome_existe": False, "driver_existe": False, "compat": False, "erro": ""}
    if not os.path.exists(CHROME_PATH):
        r["erro"] = f"Chrome não encontrado.\nCopie a pasta Bin para:\n{PASTA_BIN}"
        return r
    r["chrome_existe"] = True
    if not os.path.exists(DRIVER_PATH):
        r["erro"] = f"Chromedriver não encontrado.\nCopie a pasta Bin para:\n{PASTA_BIN}"
        return r
    r["driver_existe"] = True
    r["compat"] = True
    return r


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# REGRAS E HISTÓRICO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _para_serial_data(dt):
    return (dt.date() - EXCEL_EPOCH).days

def filtrar_modelo_inteligente(setor, modelo_site):
    setor = str(setor).upper().strip()
    modelo_site = str(modelo_site).upper().strip()

    if "ELETRO-ELETRONICOS" in setor:
        return "72X17"

    if "BELEZA" in setor:
        if "110X12" in modelo_site or "64X7" in modelo_site:
            return "64X7"
        if "68X22" in modelo_site or "46X8" in modelo_site:
            return "46X8"

    return "-"

def _atualizar_historico(ws_out, col_map, linhas_aprovadas):
    mes_atual = datetime.now().strftime("%Y-%m")
    os.makedirs(PASTA_HISTORICO, exist_ok=True)
    caminho_hist = os.path.join(PASTA_HISTORICO, f"Historico_{mes_atual}.xlsx")

    cabecalho = list(col_map.keys())

    if os.path.exists(caminho_hist):
        wb_h = load_workbook(caminho_hist)
        ws_h = wb_h.active
    else:
        wb_h = Workbook()
        ws_h = wb_h.active
        ws_h.title = "HISTORICO"
        for c_idx, nome in enumerate(cabecalho, start=1):
            ws_h.cell(row=1, column=c_idx, value=nome)

    idx_pedido = col_map.get("PEDIDO")
    idx_artigo = col_map.get("ARTIGO")
    idx_data   = col_map.get("DATA APROVADO")

    existentes = set()
    if ws_h.max_row > 1 and idx_pedido and idx_artigo and idx_data:
        for r in range(2, ws_h.max_row + 1):
            chave = (
                str(ws_h.cell(row=r, column=idx_pedido).value),
                str(ws_h.cell(row=r, column=idx_artigo).value),
                str(ws_h.cell(row=r, column=idx_data).value),
            )
            existentes.add(chave)

    proxima_linha = ws_h.max_row + 1
    adicionados = 0
    for linha_num in linhas_aprovadas:
        valores = {nome: ws_out.cell(row=linha_num, column=idx).value for nome, idx in col_map.items()}
        chave = (str(valores.get("PEDIDO")), str(valores.get("ARTIGO")), str(valores.get("DATA APROVADO")))
        if chave in existentes:
            continue
        for nome, idx in col_map.items():
            ws_h.cell(row=proxima_linha, column=idx, value=valores.get(nome))
        proxima_linha += 1
        adicionados += 1
        existentes.add(chave)

    wb_h.save(caminho_hist)
    if adicionados:
        registrar_log([f"📚  {adicionados} pedido(s) adicionados ao histórico de {mes_atual}."])

def _gravar(ws, linha, col_map, nome_coluna, valor):
    idx = col_map.get(nome_coluna)
    if not idx:
        return
    ws.cell(row=linha, column=idx, value=valor)


_PREFIXO_PERFIL_TEMP = "aprovador_chrome_perfil_"

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# NAVEGADOR
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _matar_processos_zumbis():
    for caminho in (CHROME_PATH, DRIVER_PATH):
        if not caminho or not os.path.exists(caminho):
            continue
        try:
            caminho_ps = caminho.replace("'", "''")
            ps = (
                f"Get-CimInstance Win32_Process -ErrorAction SilentlyContinue | "
                f"Where-Object {{ $_.ExecutablePath -eq '{caminho_ps}' }} | "
                f"ForEach-Object {{ Stop-Process -Id $_.ProcessId -Force -ErrorAction SilentlyContinue }}"
            )
            subprocess.run(
                ["powershell", "-NoProfile", "-WindowStyle", "Hidden", "-Command", ps],
                capture_output=True, timeout=15,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
            )
        except Exception:
            pass

def _limpar_perfis_temp_antigos(idade_min=120):
    try:
        base = tempfile.gettempdir()
        limite = time.time() - idade_min * 60
        for nome in os.listdir(base):
            if not nome.startswith(_PREFIXO_PERFIL_TEMP):
                continue
            caminho = os.path.join(base, nome)
            try:
                if os.path.getmtime(caminho) < limite:
                    shutil.rmtree(caminho, ignore_errors=True)
            except Exception:
                pass
    except Exception:
        pass

def iniciar_navegador():
    try:
        _matar_processos_zumbis()
        _limpar_perfis_temp_antigos()
        options = Options()
        options.binary_location = CHROME_PATH
        options.add_argument("--start-maximized")
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")


        perfil_temp = tempfile.mkdtemp(prefix=f"{_PREFIXO_PERFIL_TEMP}{os.getpid()}_")
        options.add_argument(f"--user-data-dir={perfil_temp}")


        options.add_argument("--no-first-run")
        options.add_argument("--no-default-browser-check")
        options.add_argument("--disable-notifications")
        options.add_argument("--disable-popup-blocking")
        options.add_experimental_option("prefs", {
            "credentials_enable_service": False,
            "profile.password_manager_enabled": False,
        })
        options.add_experimental_option("excludeSwitches", ["enable-automation"])
        servico = Service(executable_path=DRIVER_PATH)
        driver = webdriver.Chrome(service=servico, options=options)


        driver.set_page_load_timeout(180)
        return driver
    except Exception as e:
        registrar_log([f"❌  Erro ao abrir navegador: {e}"])
        registrar_erro_permanente(f"NAVEGADOR -- {e}")
        return None

_ultimo_erro_selenium = ""


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# AÇÕES NA PÁGINA (SELENIUM)
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def esperar(driver, xpath, condicao="visible", timeout=10):
    global _ultimo_erro_selenium
    try:
        wait = WebDriverWait(driver, timeout, poll_frequency=0.2)
        if condicao == "clickable":
            return wait.until(EC.element_to_be_clickable((By.XPATH, xpath)))
        if condicao == "visible":
            return wait.until(EC.visibility_of_element_located((By.XPATH, xpath)))
        return wait.until(EC.presence_of_element_located((By.XPATH, xpath)))
    except Exception as e:
        _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
        return None

def _preencher_campo(driver, elemento, valor):
    global _ultimo_erro_selenium
    try:
        elemento.clear()
        elemento.send_keys(valor)
        valor_atual = elemento.get_attribute("value")
        if valor_atual == valor:
            return True
        _ultimo_erro_selenium = (f"clear()+send_keys() não substituiu direito -- "
                                 f"campo ficou com {valor_atual!r} em vez de {valor!r}")
    except Exception as e:
        _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"

    try:
        driver.execute_script("""
            const el = arguments[0];
            const val = arguments[1];
            const setter = Object.getOwnPropertyDescriptor(window.HTMLInputElement.prototype, 'value').set;
            setter.call(el, val);
            el.dispatchEvent(new Event('input', { bubbles: true }));
            el.dispatchEvent(new Event('change', { bubbles: true }));
        """, elemento, valor)


        if (elemento.get_attribute("value") or "") == valor:


            _ultimo_erro_selenium = ""
            return True
        _ultimo_erro_selenium = (f"JS setter também não substituiu direito -- "
                                 f"campo ficou com {(elemento.get_attribute('value') or '')!r} em vez de {valor!r}")
        return False
    except Exception as e:
        _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
        return False

_JS_CLICAR_XPATH = """
    const xpath = arguments[0];
    const el = document.evaluate(xpath, document, null, XPathResult.FIRST_ORDERED_NODE_TYPE, null).singleNodeValue;
    if (!el) return 'nao_encontrado';
    const visivel = el.getClientRects().length > 0 && el.offsetParent !== null;
    if (!visivel) return 'nao_visivel';
    if (el.disabled) return 'desabilitado';
    el.scrollIntoView({block: 'nearest', inline: 'center'});
    el.click();
    return 'clicado';
"""

# Clique via JS em loop > resiste a overlay e animação; _clicar_nativo é o clique Selenium padrão.
def clicar(driver, xpath, timeout=6):
    global _ultimo_erro_selenium
    intervalo = 0.15
    tentativas = max(1, int(timeout / intervalo))
    resultado = "nao_encontrado"
    for _ in range(tentativas):
        try:
            resultado = driver.execute_script(_JS_CLICAR_XPATH, xpath)
        except Exception as e:
            _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
            return False
        if resultado == "clicado":
            return True
        time.sleep(intervalo)
    _ultimo_erro_selenium = f"elemento não ficou clicável a tempo (último estado: {resultado})"
    return False

def _clicar_nativo(driver, xpath, timeout=5):
    global _ultimo_erro_selenium
    if not xpath:
        return False
    try:
        el = WebDriverWait(driver, timeout).until(EC.element_to_be_clickable((By.XPATH, xpath)))
        el.click()
        return True
    except Exception as e:
        _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
        return False

def _texto_celula(driver, indice_linha, chave_sufixo):
    xpath = f"{XPATHS['linhas_tabela_pedido']}[{indice_linha}]/{XPATHS[chave_sufixo]}"
    el = esperar(driver, xpath, "presence", 5)
    return el.text.strip() if el else ""

_MAPA_COLUNAS_DETALHE_PEDIDO = None

def _resetar_mapa_colunas_detalhe_pedido():
    global _MAPA_COLUNAS_DETALHE_PEDIDO
    _MAPA_COLUNAS_DETALHE_PEDIDO = None

def _mapa_colunas_detalhe_pedido(driver, forcar=False):
    global _MAPA_COLUNAS_DETALHE_PEDIDO
    if _MAPA_COLUNAS_DETALHE_PEDIDO is not None and not forcar:
        return _MAPA_COLUNAS_DETALHE_PEDIDO
    try:
        linha = esperar(driver, f"{XPATHS['linhas_tabela_pedido']}[1]", "presence", 3)
        if not linha:
            return {}
        nomes = driver.execute_script(
            """
            const linha = arguments[0];
            const tabela = linha.closest('table');
            if (!tabela) return [];
            const thead = tabela.querySelector('thead');
            if (!thead) return [];
            return Array.from(thead.querySelectorAll('th')).map(th => th.textContent.trim());
            """,
            linha,
        )
        mapa = {nome.strip().lower(): i + 1 for i, nome in enumerate(nomes or []) if nome and nome.strip()}
        _MAPA_COLUNAS_DETALHE_PEDIDO = mapa
        return mapa
    except Exception:
        return {}

def _texto_celula_por_nome(driver, indice_linha, nome_coluna, chave_sufixo_fallback):
    mapa = _mapa_colunas_detalhe_pedido(driver)
    indice_col = mapa.get(nome_coluna.strip().lower())
    if indice_col:
        xpath = f"{XPATHS['linhas_tabela_pedido']}[{indice_linha}]/td[{indice_col}]"
        el = esperar(driver, xpath, "presence", 5)
        if el:
            return el.text.strip()
    return _texto_celula(driver, indice_linha, chave_sufixo_fallback)

def _validar_artigo(driver, site_sku, planilha_sku):
    if site_sku != planilha_sku:
        time.sleep(0.5)
        site_sku_confere = _texto_celula_por_nome(driver, 1, "SKU", "sufixo_sku")[:6]
        if site_sku_confere == planilha_sku:
            site_sku = site_sku_confere
    res_artigo = "OK" if site_sku == planilha_sku else "DIVERGENTE"
    return res_artigo, site_sku

def _xpath_botao_editar_qtd(driver, indice_linha):
    mapa = _mapa_colunas_detalhe_pedido(driver)
    indice_col = mapa.get("ação") or mapa.get("acao")
    if indice_col:
        return f"{XPATHS['linhas_tabela_pedido']}[{indice_linha}]/td[{indice_col}]/div/button[2]"
    return _xpath_celula(indice_linha, "sufixo_btn_alterar_qtd")

def _texto_xpath(driver, chave, timeout=5):
    el = esperar(driver, XPATHS.get(chave, ""), "presence", timeout)
    return el.text.strip() if el else ""

def _xpath_celula(indice_linha, chave_sufixo):
    return f"{XPATHS['linhas_tabela_pedido']}[{indice_linha}]/{XPATHS[chave_sufixo]}"

def _limpar_filtro(driver):
    clicar(driver, XPATHS["btn_limpar_filtro_pedido"], timeout=4)
    time.sleep(0.5)

def _selecionar_opcao_dropdown(driver, texto_parcial, timeout=8):
    global _ultimo_erro_selenium
    texto_lower = texto_parcial.lower()
    fim = time.time() + timeout
    while time.time() < fim:
        try:
            candidatos = driver.find_elements(By.XPATH, "//div[@role='option'] | //li")
            for el in candidatos:
                try:
                    texto_el = driver.execute_script("return arguments[0].textContent || '';", el)
                    if texto_lower in texto_el.strip().lower():
                        return el
                except Exception as e:
                    _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
                    continue
        except Exception as e:
            _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
        time.sleep(0.3)
    if not _ultimo_erro_selenium:
        _ultimo_erro_selenium = "Nenhuma opção com esse texto foi encontrada na lista (sem erro de Selenium)."
    return None

def _clicar_elemento(driver, elemento):
    global _ultimo_erro_selenium
    try:
        elemento.click()
        return True
    except Exception:
        try:
            driver.execute_script("arguments[0].click();", elemento)
            return True
        except Exception as e:
            _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
            return False

def _esperar_carregamento_sumir(driver, timeout=15):
    try:
        WebDriverWait(driver, timeout).until_not(
            EC.visibility_of_any_elements_located(
                (By.CSS_SELECTOR, ".v-overlay--active, .v-progress-circular, .v-overlay__scrim")
            )
        )
    except Exception:
        pass

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# DIAGNÓSTICO DE ERRO E HISTÓRICO PARA IMPRESSÃO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _pasta_erro_do_dia(agora=None):
    agora = agora or datetime.now()
    return os.path.join(_pasta_logs, "ERRO", agora.strftime("%Y"), agora.strftime("%m"), agora.strftime("%d.%m.%y"))

def _proximo_codigo_erro():
    agora = datetime.now()
    pasta_dia = _pasta_erro_do_dia(agora)
    try:
        os.makedirs(pasta_dia, exist_ok=True)
        existentes = [n for n in os.listdir(pasta_dia) if n.startswith("ERRO")]
        numeros = []
        for n in existentes:
            try:
                numeros.append(int(n.replace("ERRO", "")))
            except Exception:
                continue
        proximo = (max(numeros) + 1) if numeros else 1
    except Exception:
        proximo = 1
    nome_pasta = f"ERRO{proximo:03d}"
    codigo_legivel = f"{agora.strftime('%d.%m.%y')}_{nome_pasta}"
    return nome_pasta, codigo_legivel

def _pasta_historico_impressao(departamento, data_dt):
    departamento = str(departamento or "SEM_SETOR").strip().upper()
    return os.path.join(PASTA_HISTORICO_IMPRESSAO, departamento,
                         data_dt.strftime("%Y"), data_dt.strftime("%m"), data_dt.strftime("%d"))

def _proxima_versao_historico(pasta_dia):
    try:
        os.makedirs(pasta_dia, exist_ok=True)
        numeros = []
        for nome in os.listdir(pasta_dia):
            m = re.match(r"[Vv](\d+)\.xlsx$", nome)
            if m:
                numeros.append(int(m.group(1)))
        return (max(numeros) + 1) if numeros else 1
    except Exception:
        return 1

def _salvar_historico_impressao(wb_in, wb_out):
    try:
        if "CONFIGURAÇÕES" not in wb_in.sheetnames:
            registrar_log(["⚠️  Aba CONFIGURAÇÕES não encontrada -- cópia para o Robô de Impressão não salva."])
            return
        ws_cfg = wb_in["CONFIGURAÇÕES"]
        data_valor   = ws_cfg["A2"].value
        departamento = ws_cfg["B2"].value

        if isinstance(data_valor, datetime):
            data_dt = data_valor
        elif isinstance(data_valor, ddate):
            data_dt = datetime(data_valor.year, data_valor.month, data_valor.day)
        else:
            data_dt = datetime.strptime(str(data_valor).strip(), "%d/%m/%Y")

        pasta_dia = _pasta_historico_impressao(departamento, data_dt)
        versao = _proxima_versao_historico(pasta_dia)
        caminho_versao = os.path.join(pasta_dia, f"V{versao}.xlsx")
        wb_out.save(caminho_versao)
        registrar_log([f"🖨️  Cópia salva para o Robô de Impressão: {caminho_versao}"])
    except Exception as e:
        registrar_log([f"⚠️  Falha ao salvar cópia para o Robô de Impressão: {type(e).__name__}: {e}"])
        registrar_erro_permanente(f"HISTÓRICO IMPRESSÃO -- {type(e).__name__}: {e}")

def _listar_campos_da_pagina(driver):
    linhas = []
    try:
        driver.switch_to.default_content()
        contextos = [("documento principal", None)]
        try:
            iframes = driver.find_elements(By.TAG_NAME, "iframe")
            for idx in range(len(iframes)):
                contextos.append((f"iframe {idx}", idx))
        except Exception:
            pass

        for nome_ctx, idx_iframe in contextos:
            try:
                driver.switch_to.default_content()
                if idx_iframe is not None:
                    iframes_atual = driver.find_elements(By.TAG_NAME, "iframe")
                    if idx_iframe >= len(iframes_atual):
                        continue
                    driver.switch_to.frame(iframes_atual[idx_iframe])
                campos = driver.find_elements(By.XPATH, "//input[@type='text' or @type='password' or not(@type)]")
                if not campos:
                    continue
                linhas.append(f"[{nome_ctx}] {len(campos)} campo(s) de input encontrado(s):")
                for c in campos:
                    try:
                        cid = c.get_attribute("id") or "(sem id)"
                        ctipo = c.get_attribute("type") or "(sem type)"
                        cvis = "visível" if c.is_displayed() else "oculto"
                        linhas.append(f"   id={cid}  type={ctipo}  {cvis}")
                    except Exception:
                        continue
            except Exception:
                continue
    except Exception:
        pass
    finally:
        try:
            driver.switch_to.default_content()
        except Exception:
            pass
    return linhas or ["Nenhum campo de input encontrado em lugar nenhum."]

def _diagnostico_pagina(driver):
    try:
        return driver.current_url, driver.title
    except Exception:
        return "?", "?"

def _capturar_diagnostico(driver, titulo, com_html=False, com_campos=False):
    nome_pasta, codigo = _proximo_codigo_erro()
    pasta = os.path.join(_pasta_erro_do_dia(), nome_pasta)
    try:
        os.makedirs(pasta, exist_ok=True)
    except Exception:
        pass

    try:
        driver.save_screenshot(os.path.join(pasta, "print.png"))
    except Exception:
        pass


    try:
        with open(os.path.join(pasta, "xpaths_carregados.json"), "w", encoding="utf-8") as f:
            json.dump(XPATHS, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

    if com_html:
        try:
            with open(os.path.join(pasta, "pagina.html"), "w", encoding="utf-8") as f:
                f.write(driver.page_source)
        except Exception:
            pass

    url_atual, titulo_pagina = _diagnostico_pagina(driver)
    linhas = [
        f"Código do erro: {codigo}",
        titulo,
        f"URL: {url_atual}",
        f"Título da página: {titulo_pagina}",
        f"Erro real do Selenium: {_ultimo_erro_selenium}",
        "(xpaths_aprovador.json carregado na memória nesse instante salvo em xpaths_carregados.json, nesta mesma pasta)",
    ]
    if com_campos:
        linhas.append("Campos encontrados na página:")
        linhas.extend(_listar_campos_da_pagina(driver))

    try:
        with open(os.path.join(pasta, "log.txt"), "w", encoding="utf-8") as f:
            f.write("\n".join(linhas))
    except Exception:
        pass

    registrar_log(linhas + [f"🗂️  Diagnóstico completo em: {pasta}"])
    return codigo

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# BUSCA DE CAMPOS (COM FALLBACK)
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _procurar_em_iframes(driver, xpath, timeout=8, condicao="presence"):
    driver.switch_to.default_content()
    elemento = esperar(driver, xpath, condicao, timeout)
    if elemento:
        return elemento, None

    try:
        iframes = driver.find_elements(By.TAG_NAME, "iframe")
    except Exception:
        iframes = []

    for idx, iframe in enumerate(iframes):
        try:
            driver.switch_to.default_content()
            driver.switch_to.frame(iframe)
            elemento = esperar(driver, xpath, condicao, 5)
            if elemento:
                return elemento, idx
        except Exception:
            continue

    driver.switch_to.default_content()
    return None, None

# XPath configurado primeiro, depois o de reserva > o aviso indica que o principal precisa de revisão.
def _procurar_campo(driver, chave_principal, chave_fallback=None, timeout=10,
                    condicao="presence", usar_iframes=False):
    xpath_principal = XPATHS.get(chave_principal, "")
    if usar_iframes:
        elemento, iframe_idx = _procurar_em_iframes(driver, xpath_principal, timeout, condicao)
    else:
        elemento = esperar(driver, xpath_principal, condicao, timeout)
        iframe_idx = None
    if elemento:
        return elemento, iframe_idx, False

    xpath_fallback = XPATHS.get(chave_fallback or "", "")
    if not xpath_fallback:
        return None, None, False

    if usar_iframes:
        elemento, iframe_idx = _procurar_em_iframes(driver, xpath_fallback, timeout, condicao)
    else:
        elemento = esperar(driver, xpath_fallback, condicao, timeout)
        iframe_idx = None
    if elemento:
        registrar_log([f"⚠️  '{chave_principal}' não encontrado pelo XPath configurado -- "
                       f"XPath de reserva (por placeholder) utilizado; revisar "
                       f"'{chave_principal}' no xpaths_aprovador.json."])
        return elemento, iframe_idx, True

    return None, None, False

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LOGIN E NAVEGAÇÃO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def fazer_login(driver):
    for tentativa in (1, 2):
        try:
            driver.get(PORTAL_URL_LOGIN)
            timeout_campo = 20 if tentativa == 1 else 35
            _esperar_carregamento_sumir(driver, timeout=timeout_campo)

            campo_user, iframe_idx, _ = _procurar_campo(
                driver, "campo_usuario", "campo_usuario_fallback",
                timeout_campo, "presence", usar_iframes=True)
            if not campo_user:
                _capturar_diagnostico(
                    driver,
                    f"⚠️  Tentativa {tentativa}/2: campo de usuário não apareceu (nem no documento principal, nem em iframes)",
                    com_html=(tentativa == 2), com_campos=(tentativa == 2)
                )
                if tentativa == 1:
                    time.sleep(3)
                    continue
                return False, "Campo de usuário não apareceu (após 2 tentativas)"
            if iframe_idx is not None:
                registrar_log([f"ℹ️  Campo de login encontrado dentro de um iframe (índice {iframe_idx})."])
            if not _preencher_campo(driver, campo_user, PORTAL_USER):
                registrar_log([f"⚠️  Tentativa {tentativa}/2: falha ao preencher o campo de usuário"])
                if tentativa == 1:
                    time.sleep(3)
                    continue
                return False, "Falha ao preencher o campo de usuário"


            campo_senha, _, _ = _procurar_campo(driver, "campo_senha", "campo_senha_fallback", 10, "presence")
            if not campo_senha:
                _capturar_diagnostico(driver, f"⚠️  Tentativa {tentativa}/2: campo de senha não apareceu")
                if tentativa == 1:
                    time.sleep(3)
                    continue
                return False, "Campo de senha não apareceu (após 2 tentativas)"
            if not _preencher_campo(driver, campo_senha, PORTAL_SENHA):
                registrar_log([f"⚠️  Tentativa {tentativa}/2: falha ao preencher o campo de senha"])
                if tentativa == 1:
                    time.sleep(3)
                    continue
                return False, "Falha ao preencher o campo de senha"

            clicar(driver, XPATHS["btn_login"])
            time.sleep(3)


            driver.switch_to.default_content()
            _esperar_carregamento_sumir(driver, timeout=15)


            campo_site, _, _ = _procurar_campo(driver, "campo_site", "campo_site_fallback", 15, "presence")
            if not campo_site:
                _capturar_diagnostico(
                    driver, f"⚠️  Tentativa {tentativa}/2: campo de Site não apareceu",
                    com_html=(tentativa == 2), com_campos=(tentativa == 2)
                )
                if tentativa == 1:
                    time.sleep(3)
                    continue
                return False, "Campo de Site não apareceu (após 2 tentativas)"

            _preencher_campo(driver, campo_site, PORTAL_SITE)
            time.sleep(1.5)
            opcao = _selecionar_opcao_dropdown(driver, PORTAL_SITE.lower())


            if not opcao or not _clicar_elemento(driver, opcao):
                _capturar_diagnostico(
                    driver, f"⚠️  Tentativa {tentativa}/2: falha ao selecionar '{PORTAL_SITE}' na lista de sites"
                )
                if tentativa == 1:
                    time.sleep(3)
                    continue
                return False, f"Falha ao selecionar o site '{PORTAL_SITE}' (após 2 tentativas)"
            time.sleep(0.5)

            clicar(driver, XPATHS["btn_entrar_site"])
            time.sleep(3)
            return True, ""

        except Exception as e:
            try:
                driver.switch_to.default_content()
            except Exception:
                pass
            if tentativa == 1:
                registrar_log([f"⚠️  Tentativa 1/2 falhou com erro inesperado: "
                              f"{type(e).__name__}: {e}. Tentando de novo..."])
                time.sleep(3)
                continue
            _capturar_diagnostico(driver, f"❌  Falha inesperada no login: {type(e).__name__}: {e}")
            return False, str(e)

    return False, "Falha desconhecida no login"

def navegar_para_pedidos(driver):
    if not clicar(driver, XPATHS["menu_grupo_pedido"], timeout=8):
        codigo = _capturar_diagnostico(
            driver, "❌  Falha ao clicar no menu 'Pedido' (grupo do menu não ficou clicável).",
            com_html=True
        )
        registrar_erro_permanente(f"NAVEGAÇÃO PEDIDOS -- menu 'Pedido' não clicável (código {codigo})")
        return False
    time.sleep(1)

    if not clicar(driver, XPATHS["menu_item_pedido"], timeout=8):
        codigo = _capturar_diagnostico(
            driver, "❌  Grupo 'Pedido' clicado, mas o item do menu não ficou clicável.",
            com_html=True
        )
        registrar_erro_permanente(f"NAVEGAÇÃO PEDIDOS -- item do menu não clicável (código {codigo})")
        return False
    time.sleep(2)


    confirmou, _, _ = _procurar_campo(driver, "campo_filtrar_pedido_fallback", "campo_filtrar_pedido", 8, "presence")
    if confirmou:


        fim_poll = time.time() + 5
        while time.time() < fim_poll:
            try:
                if driver.find_elements(By.XPATH, XPATHS["linhas_tabela_pedido"]):
                    break
            except Exception:
                pass
            time.sleep(0.3)


        time.sleep(3)
        return True

    codigo = _capturar_diagnostico(
        driver, "❌  Itens do menu clicados, mas a tela de Pedidos não foi confirmada.",
        com_html=True
    )
    registrar_erro_permanente(f"NAVEGAÇÃO PEDIDOS -- não confirmou a tela certa (código {codigo})")
    return False

class ParadaCritica(Exception):
    pass

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# FILTRO DE SITE
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _limpar_filtro_periodo(driver):
    try:
        campos = driver.find_elements(By.XPATH, XPATHS.get("campo_periodo", ""))
        if not campos:
            return
        campo = campos[0]
        if not campo.get_attribute("value"):
            return
        campo.click()
        campo.send_keys(Keys.CONTROL + "a")
        campo.send_keys(Keys.DELETE)
        campo.send_keys(Keys.ENTER)


        campo.send_keys(Keys.ESCAPE)
        time.sleep(0.5)
    except Exception:
        pass

_JS_STATUS_FILTRO_SITE = """
    const containers = document.querySelectorAll('.v-input');
    for (const cont of containers) {
        const label = cont.querySelector('label.v-field-label, label.v-label');
        if (!label || label.textContent.trim() !== 'Site') continue;
        const selecao = cont.querySelector('.v-autocomplete__selection .truncate');
        const texto = selecao ? selecao.textContent.trim() : '';
        if (texto && !['todos', 'todos os sites', 'all'].includes(texto.toLowerCase())) {
            return {precisa: true, texto: texto};
        }
    }
    return {precisa: false, texto: ''};
"""

_JS_LIMPAR_FILTRO_SITE = """
    const containers = document.querySelectorAll('.v-input');
    for (const cont of containers) {
        const label = cont.querySelector('label.v-field-label, label.v-label');
        if (!label || label.textContent.trim() !== 'Site') continue;
        const selecao = cont.querySelector('.v-autocomplete__selection .truncate');
        const texto = selecao ? selecao.textContent.trim() : '';
        if (texto && !['todos', 'todos os sites', 'all'].includes(texto.toLowerCase())) {
            const icone = cont.querySelector("i[aria-label='Clear Site']");
            if (!icone) return 'sem_icone';
            icone.scrollIntoView({block: 'nearest', inline: 'center'});
            icone.click();
            return 'clicado';
        }
    }
    return 'nao_precisava';
"""

def _filtro_ainda_mostra_site(driver):
    try:
        resultado = driver.execute_script(_JS_STATUS_FILTRO_SITE)
        return bool(resultado and resultado.get("precisa"))
    except Exception:
        return False

def _filtro_precisa_resetar(driver):
    if _filtro_ainda_mostra_site(driver):
        return True
    time.sleep(0.4)
    return _filtro_ainda_mostra_site(driver)

def _esperar_filtro_resetar(driver, timeout=1.5, intervalo=0.2):
    fim = time.time() + timeout
    while time.time() < fim:
        if not _filtro_ainda_mostra_site(driver):
            return False
        time.sleep(intervalo)
    return _filtro_ainda_mostra_site(driver)

def _resetar_filtro_site(driver, pedido):


    if not _filtro_precisa_resetar(driver):
        return


    if _clicar_nativo(driver, XPATHS["campo_filtro_site_pedidos"], timeout=4):
        opcao = None
        for texto_candidato in ("todos", "todos os sites", "all"):
            opcao = _selecionar_opcao_dropdown(driver, texto_candidato, timeout=3)
            if opcao:
                break
        if opcao and _clicar_elemento(driver, opcao):
            if not _esperar_filtro_resetar(driver):
                return


    try:
        driver.execute_script(_JS_LIMPAR_FILTRO_SITE)
    except Exception:
        pass
    if not _esperar_filtro_resetar(driver):
        return


    codigo_erro = _capturar_diagnostico(
        driver,
        f"⚠️  Filtro de Site não resetou (nem pela lista, nem pelo 'X') antes de {pedido}",
        com_html=True
    )
    raise ParadaCritica(
        f"Filtro de Site não resetou antes de {pedido} (diagnóstico: {codigo_erro})"
    )

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# SELECIONAR PEDIDO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_JS_ACHAR_ROLAR_CLICAR_OLHO = """
    const pedido = arguments[0];
    const classeIcone = arguments[1];
    const linhas = document.querySelectorAll('table tbody tr');
    for (const tr of linhas) {
        const tds = tr.querySelectorAll('td');
        let bate = false;
        for (const td of tds) {
            if (td.textContent.trim() === pedido) { bate = true; break; }
        }
        if (bate) {
            const icone = tr.querySelector(`button i.${classeIcone}, button i[class*='${classeIcone}']`);
            if (!icone) {
                return 'linha_sem_botao';
            }
            const botao = icone.closest('button');
            botao.scrollIntoView({block: 'nearest', inline: 'center'});
            botao.click();
            return 'clicado';
        }
    }
    return 'linha_nao_encontrada';
"""

_XPATH_NUMERO_PEDIDO = "//span[starts-with(normalize-space(.),'Número do Pedido:')]"

def _esperar_detalhe_do_pedido(driver, pedido, timeout=10):
    alvo = str(pedido).strip().lstrip("0")
    fim = time.time() + timeout
    while time.time() < fim:
        try:
            for el in driver.find_elements(By.XPATH, _XPATH_NUMERO_PEDIDO):
                numero = el.text.split(":", 1)[-1].strip().lstrip("0")
                if numero and numero == alvo:
                    return True
        except Exception:
            pass
        time.sleep(0.3)
    return False

def _tentar_clicar_olho(driver, pedido, tentativas=20, intervalo=0.5):
    global _ultimo_erro_selenium
    classe_icone = XPATHS.get("classe_icone_olho", "mdi-eye-outline")
    motivo = "linha_nao_encontrada"
    for _ in range(tentativas):
        try:
            motivo = driver.execute_script(_JS_ACHAR_ROLAR_CLICAR_OLHO, pedido, classe_icone)
        except Exception as e:
            _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
            return False, "erro_js"
        if motivo == "clicado":
            return True, motivo
        if motivo == "linha_sem_botao":
            break
        time.sleep(intervalo)
    return False, motivo

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# DIÁLOGOS DE CONFIRMAÇÃO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_JS_CLICAR_CONFIRMACAO_DIALOGO = """
    const seletores = ['.v-overlay--active .v-card', '.v-dialog--active .v-card', "[role='dialog']"];
    for (const sel of seletores) {
        const cards = document.querySelectorAll(sel);
        for (const card of cards) {
            const botoes = Array.from(card.querySelectorAll('button')).filter(b => b.offsetParent !== null);
            if (botoes.length > 0) {
                const alvo = botoes[botoes.length - 1];
                alvo.scrollIntoView({block: 'nearest'});
                alvo.click();
                return 'clicado';
            }
        }
    }
    return 'dialogo_nao_encontrado';
"""

def _clicar_confirmacao_dialogo(driver, xpath_fragil, tentativas=6, intervalo=0.5):
    if xpath_fragil and clicar(driver, xpath_fragil, timeout=5):
        return True

    global _ultimo_erro_selenium
    for _ in range(tentativas):
        try:
            resultado = driver.execute_script(_JS_CLICAR_CONFIRMACAO_DIALOGO)
        except Exception as e:
            _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
            resultado = "erro_js"
        if resultado == "clicado":


            registrar_log([
                "ℹ️  Diálogo de confirmação: XPath configurado não funcionou, "
                "clique de reserva (JS) utilizado com sucesso; revisar o XPath principal."
            ])
            return True
        time.sleep(intervalo)
    return False

_JS_ACHAR_CAMPO_QTD_DIALOGO = """
    const seletores = ['.v-overlay--active .v-card', '.v-dialog--active .v-card', "[role='dialog']"];
    for (const sel of seletores) {
        const cards = document.querySelectorAll(sel);
        for (const card of cards) {
            const inputs = Array.from(card.querySelectorAll('input')).filter(i => i.offsetParent !== null);
            if (inputs.length > 0) return inputs[0];
        }
    }
    return null;
"""

def _achar_campo_qtd_dialogo(driver, tentativas=18, intervalo=0.3):
    global _ultimo_erro_selenium
    for _ in range(tentativas):
        try:
            el = driver.execute_script(_JS_ACHAR_CAMPO_QTD_DIALOGO)
        except Exception as e:


            _ultimo_erro_selenium = f"{type(e).__name__}: {str(e)[:300]}"
            el = None
        if el:
            return el
        time.sleep(intervalo)


    campo = esperar(driver, XPATHS.get("campo_editar_qtd", ""), "visible", 2)
    if campo:
        registrar_log(["⚠️  Campo de editar Qtd: não localizado pelo diálogo ativo -- "
                       "ID fixo de reserva ('campo_editar_qtd') utilizado; revisar o XPath principal."])
    return campo

def _fechar_dialogo_qtd(driver):
    if clicar(driver, XPATHS.get("btn_cancelar_qtd_fragil", ""), timeout=3):
        return True
    if clicar(driver, XPATHS.get("btn_cancelar_dialogo", ""), timeout=3):
        registrar_log(["⚠️  Cancelar popup Qtd: XPath posicional ('btn_cancelar_qtd_fragil') não funcionou -- "
                       "XPath de reserva por texto ('btn_cancelar_dialogo') utilizado; revisar o XPath principal."])
        return True
    return False

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# PROCESSAMENTO DE PEDIDO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def processar_pedido(driver, dados, ws_out, linha_num, col_map):
    pedido = str(dados.get("PEDIDO", "") or "").strip()
    if not pedido:
        return

    local_imp = str(dados.get("LOCAL IMPRESSÃO", "") or "").strip().upper()
    if local_imp != "CDA":
        _gravar(ws_out, linha_num, col_map, "STATUS", "FORA CDA")
        return

    _resetar_mapa_colunas_detalhe_pedido()
    registrar_log([f"⚡  Processando pedido {pedido}"])


    _tempos = []
    _t_ref = [time.perf_counter()]
    def _marca(nome):
        agora = time.perf_counter()
        _tempos.append(f"{nome}={agora - _t_ref[0]:.2f}s")
        _t_ref[0] = agora

    try:


        _resetar_filtro_site(driver, pedido)
        _marca("reset_site")
        _limpar_filtro_periodo(driver)
        _marca("limpar_periodo")

        campo, _, _ = _procurar_campo(driver, "campo_filtrar_pedido_fallback", "campo_filtrar_pedido", 8, "visible")
        _marca("achar_campo_busca")
        if not campo:
            _gravar(ws_out, linha_num, col_map, "STATUS", "ERRO BUSCA")
            _capturar_diagnostico(driver, f"❌  {pedido}: campo de busca não apareceu", com_html=True)
            return
        if not _preencher_campo(driver, campo, pedido):
            _gravar(ws_out, linha_num, col_map, "STATUS", "ERRO BUSCA")
            _capturar_diagnostico(driver, f"❌  {pedido}: falha ao preencher o campo de busca", com_html=True)
            return
        _marca("preencher_busca")


        clicou, motivo = _tentar_clicar_olho(driver, pedido)
        _marca("olho_1a_tentativa")


        tentativa_num = 1
        while not clicou and motivo != "erro_js" and tentativa_num < 3:
            tentativa_num += 1
            registrar_log([f"🔄  {pedido}: tentativa {tentativa_num}/3 (motivo anterior: {motivo})..."])
            if _filtro_precisa_resetar(driver):
                _resetar_filtro_site(driver, pedido)
                _limpar_filtro_periodo(driver)
                campo, _, _ = _procurar_campo(driver, "campo_filtrar_pedido_fallback", "campo_filtrar_pedido", 8, "visible")
                if not campo:
                    continue
            if not _preencher_campo(driver, campo, pedido):
                continue
            clicou, motivo = _tentar_clicar_olho(driver, pedido)
            _marca(f"olho_tentativa_{tentativa_num}")

        if not clicou:
            _gravar(ws_out, linha_num, col_map, "STATUS", "NAO ACHOU")
            _capturar_diagnostico(
                driver,
                f"⚠️  {pedido}: não apareceu na lista após a busca ({motivo})",
                com_html=True
            )
            _limpar_filtro(driver)
            return

        if not _esperar_detalhe_do_pedido(driver, pedido):
            _gravar(ws_out, linha_num, col_map, "STATUS", "ERRO AO ABRIR PEDIDO")
            _capturar_diagnostico(
                driver,
                f"⚠️  {pedido}: a página aberta pelo ícone de detalhes não corresponde ao pedido (Número do Pedido não confere)",
                com_html=True
            )
            clicar(driver, XPATHS["btn_voltar"], timeout=5)
            _limpar_filtro(driver)
            return
        _marca("abrir_pedido")


        erro_linhas = ""
        linhas = []
        qtd_anterior = None
        fim_poll = time.time() + 3
        while time.time() < fim_poll:
            try:
                linhas = driver.find_elements(By.XPATH, XPATHS["linhas_tabela_pedido"])
            except Exception as e:
                linhas = []
                erro_linhas = f"{type(e).__name__}: {str(e)[:300]}"
            if linhas and len(linhas) == qtd_anterior:
                break
            qtd_anterior = len(linhas) if linhas else None
            time.sleep(0.2)
        qtd_linhas = len(linhas)
        _marca("leitura_tabela")

        if qtd_linhas > 1:
            _gravar(ws_out, linha_num, col_map, "STATUS", "MULTIPLAS LINHAS - REVISAR MANUAL")
            _capturar_diagnostico(
                driver,
                f"⚠️  {pedido} tem {qtd_linhas} linhas (tamanhos diferentes) -- revisão manual",
                com_html=True
            )
            clicar(driver, XPATHS["btn_voltar"], timeout=5)
            _limpar_filtro(driver)
            return

        if qtd_linhas == 0:
            _gravar(ws_out, linha_num, col_map, "STATUS", "ERRO DETALHES")
            _capturar_diagnostico(
                driver,
                f"❌  {pedido}: falha ao ler os itens do pedido -- "
                f"{erro_linhas or 'nenhuma linha encontrada (sem exceção)'}",
                com_html=True
            )
            clicar(driver, XPATHS["btn_voltar"], timeout=5)
            _limpar_filtro(driver)
            return

        site_sku    = _texto_celula_por_nome(driver, 1, "SKU", "sufixo_sku")[:6]
        site_qtd    = _texto_celula_por_nome(driver, 1, "Quantidade Total", "sufixo_qtd")
        site_modelo = _texto_celula_por_nome(driver, 1, "Etiqueta", "sufixo_etiqueta").upper()
        site_status = _texto_celula_por_nome(driver, 1, "Status", "sufixo_status").upper()
        estagio_real = _texto_xpath(driver, "estagio_pedido_selecionado").strip().upper()

        planilha_sku = str(dados.get("ARTIGO", "") or "").strip()[:6]
        setor        = str(dados.get("DEPARTAMENTO", "") or "").strip()
        modelo_final = filtrar_modelo_inteligente(setor, site_modelo)
        _marca("leitura_campos")


        if estagio_real and site_status and estagio_real != site_status:
            registrar_log([f"⚠️  {pedido}: status do item ('{site_status}') diferente do estágio real do pedido ('{estagio_real}') -- pedido duplicado no site."])
            _gravar(ws_out, linha_num, col_map, "STATUS", "PEDIDO COM STATUS DIFERENTES NO SITE, DUPLICADO")
            clicar(driver, XPATHS["btn_voltar"], timeout=5)
            _limpar_filtro(driver)
            return

        if site_status and site_status != "NOVO":


            _gravar(ws_out, linha_num, col_map, "VALIDAÇÃO MODELO", modelo_final)


            res_artigo, site_sku = _validar_artigo(driver, site_sku, planilha_sku)
            _gravar(ws_out, linha_num, col_map, "VALIDAÇÃO ARTIGO", res_artigo)
            if res_artigo == "DIVERGENTE":
                registrar_log([f"⚠️  {pedido}: já processado no site (status: {site_status}), mas artigo diverge -- "
                               f"planilha='{planilha_sku}' / site='{site_sku}'. Impressão não deve imprimir em cima disso."])
                _capturar_diagnostico(
                    driver,
                    f"⚠️  {pedido}: artigo divergente mesmo com status '{site_status}' -- "
                    f"planilha='{planilha_sku}' / site='{site_sku}'.",
                    com_html=True
                )


            try:
                v_site_ja = int(float(site_qtd.replace(",", ".")))
                v_plan_ja = int(float(str(dados.get("AGENDADO DIA", "0") or "0").replace(",", ".")))
                site_menor = bool(v_site_ja) and v_site_ja < v_plan_ja
            except Exception:
                v_site_ja, v_plan_ja, site_menor = None, None, False

            if site_menor:
                registrar_log([f"⚠️  {pedido}: já processado no site (status: {site_status}), mas o site tem MENOS "
                               f"que a planilha (site={v_site_ja} / planilha={v_plan_ja}) -- o que já existe não "
                               f"cobre, vai para Colmeia/impressão avulsa. Planilha mantida sem alteração."])
                _gravar(ws_out, linha_num, col_map, "VALIDAÇÃO QTD",
                        f"SITE ({v_site_ja}) MENOR QUE PLANILHA ({v_plan_ja})")
                _gravar(ws_out, linha_num, col_map, "STATUS", "PEDIDO JA APROVADO ANTES MAS QTD DIVERGENTE")
            elif site_status == "APROVADO":
                registrar_log([f"ℹ️  {pedido} já aprovado no site (status: {site_status})"])
                _gravar(ws_out, linha_num, col_map, "STATUS", "APROVADO")
            else:
                registrar_log([f"ℹ️  {pedido} já processado (status no site: {site_status})"])
                _gravar(ws_out, linha_num, col_map, "STATUS", "MOV")
            clicar(driver, XPATHS["btn_voltar"], timeout=5)
            _limpar_filtro(driver)
            return

        try:
            v_site = int(float(site_qtd.replace(",", ".")))
        except Exception:
            v_site = 0
        try:
            v_plan = int(float(str(dados.get("AGENDADO DIA", "0") or "0").replace(",", ".")))
        except Exception:
            v_plan = 0


        res_artigo, site_sku = _validar_artigo(driver, site_sku, planilha_sku)
        res_qtd      = "OK"
        status_final = ""
        pode_aprovar = False

        if res_artigo == "OK":
            if v_site != v_plan:
                if v_plan > v_site:
                    registrar_log([f"✏️   {pedido}: planilha ({v_plan}) > site ({v_site}). Editando no portal..."])
                    if clicar(driver, _xpath_botao_editar_qtd(driver, 1), timeout=5):
                        campo_qtd = _achar_campo_qtd_dialogo(driver)


                        if campo_qtd and _preencher_campo(driver, campo_qtd, str(v_plan)):
                            confirmou_qtd = (
                                _clicar_nativo(driver, XPATHS.get("btn_confirmar_qtd_fragil", ""), timeout=5)
                                or _clicar_confirmacao_dialogo(driver, XPATHS.get("btn_confirmar_qtd_fragil", ""))
                            )


                            if confirmou_qtd:
                                _clicar_confirmacao_dialogo(driver, XPATHS.get("btn_confirmar_ajuste_qtd", ""))


                                qtd_pos_confirmar = ""
                                fim_poll_qtd = time.time() + 4
                                while time.time() < fim_poll_qtd:
                                    qtd_pos_confirmar = _texto_celula_por_nome(driver, 1, "Quantidade Total", "sufixo_qtd")
                                    try:
                                        if int(float(qtd_pos_confirmar.replace(",", "."))) == v_plan:
                                            break
                                    except Exception:
                                        pass
                                    time.sleep(0.3)
                                try:
                                    qtd_bateu = int(float(qtd_pos_confirmar.replace(",", "."))) == v_plan
                                except Exception:
                                    qtd_bateu = False
                                if qtd_bateu:
                                    status_final = "APROVADO COM QTD DO SITE ALTERADA"
                                    res_qtd = "ALTERADA NO SITE"
                                    pode_aprovar = True
                                else:
                                    status_final = "PEDIDO CORRETO, ERRO AO ALTERAR QTD, VERIFICAR"
                                    res_qtd = f"TENTOU ALTERAR PARA {v_plan}, SITE FICOU COM '{qtd_pos_confirmar}'"
                                    _capturar_diagnostico(
                                        driver,
                                        f"❌  {pedido}: alteração de quantidade confirmada, mas o site continua "
                                        f"mostrando '{qtd_pos_confirmar}' em vez de '{v_plan}' -- pedido não aprovado.",
                                        com_html=True
                                    )


                                    _fechar_dialogo_qtd(driver)
                            else:
                                status_final = "ERRO EDITAR QTD"
                                _capturar_diagnostico(
                                    driver,
                                    f"❌  {pedido}: quantidade alterada, mas sem confirmação -- fechando o popup",
                                    com_html=True
                                )
                                _fechar_dialogo_qtd(driver)
                        else:
                            status_final = "ERRO EDITAR QTD"
                            _capturar_diagnostico(
                                driver,
                                f"❌  {pedido}: "
                                + ("campo de editar quantidade não apareceu" if not campo_qtd
                                   else "falha ao digitar a nova quantidade") + " -- fechando o popup",
                                com_html=True
                            )
                            _fechar_dialogo_qtd(driver)
                    else:
                        status_final = "ERRO LAPIS"
                        _capturar_diagnostico(
                            driver, f"❌  {pedido}: falha ao clicar no botão de alterar quantidade", com_html=True
                        )
                else:
                    registrar_log([f"🔄  {pedido}: site ({v_site}) > planilha ({v_plan}). Atualizando planilha..."])
                    _gravar(ws_out, linha_num, col_map, "AGENDADO DIA", v_site)
                    status_final = "APROVADO COM QTD ALTERADA"
                    res_qtd = "QTD ALTERADA NA PLANILHA"
                    pode_aprovar = True
            else:
                status_final = "APROVADO"
                pode_aprovar = True
        else:
            res_qtd = "DIVERGENTE"
            status_final = "REPROVADO"


            _capturar_diagnostico(
                driver,
                f"⛔  {pedido}: artigo divergente -- planilha='{planilha_sku}' / site='{site_sku}' -- REPROVADO.",
                com_html=True
            )

        if pode_aprovar:
            if clicar(driver, XPATHS["btn_aprovar"], timeout=5):
                _clicar_confirmacao_dialogo(driver, XPATHS.get("btn_confirmar_aprovacao_fragil", ""))


                status_recheck = ""
                fim_poll = time.time() + 3
                while time.time() < fim_poll:
                    status_recheck = _texto_celula_por_nome(driver, 1, "Status", "sufixo_status").upper()
                    if status_recheck == "APROVADO":
                        break
                    time.sleep(0.2)
                if status_recheck == "APROVADO":
                    registrar_log([f"🎉  {pedido} aprovado! Status: {status_final}"])
                    data_txt = None
                    try:
                        data_el = esperar(driver, XPATHS["data_aprovacao_span"], "presence", 3)
                        data_txt = data_el.text.strip() if data_el else None
                    except Exception:
                        data_txt = None
                    if data_txt:
                        try:
                            dt = datetime.strptime(data_txt[:10], "%d/%m/%Y")
                            _gravar(ws_out, linha_num, col_map, "DATA APROVADO", _para_serial_data(dt))
                        except Exception:
                            pass
                else:


                    status_final = "PEDIDO CORRETO, ERRO AO APROVAR, VERIFICAR"
                    _capturar_diagnostico(
                        driver, f"❌  {pedido}: status não mudou para APROVADO após confirmar", com_html=True
                    )
            else:


                estagio_real = _texto_xpath(driver, "estagio_pedido_selecionado").strip().upper()
                if estagio_real and estagio_real not in ("NOVO", "APROVADO"):
                    registrar_log([f"ℹ️  {pedido}: item aparecia como NOVO mas o pedido já está em '{estagio_real}' -- status do item no site está desatualizado."])
                    status_final = "MOV"
                else:
                    status_final = "ERRO APROVAR"
                    _capturar_diagnostico(
                        driver, f"❌  {pedido}: botão Aprovar não clicável", com_html=True
                    )
            _marca("aprovar")

        _gravar(ws_out, linha_num, col_map, "VALIDAÇÃO QTD", res_qtd)
        _gravar(ws_out, linha_num, col_map, "VALIDAÇÃO ARTIGO", res_artigo)
        _gravar(ws_out, linha_num, col_map, "VALIDAÇÃO MODELO", modelo_final)
        _gravar(ws_out, linha_num, col_map, "STATUS", status_final)

        clicar(driver, XPATHS["btn_voltar"], timeout=5)
        _limpar_filtro(driver)
        _marca("voltar")

    except ParadaCritica:
        raise
    except WebDriverException:
        raise

    except Exception as e:
        registrar_log([f"❌  Erro inesperado no pedido {pedido}: {type(e).__name__}: {e}"])
        _gravar(ws_out, linha_num, col_map, "STATUS", "ERRO")
        registrar_erro_permanente(f"PEDIDO {pedido} -- {type(e).__name__}: {e}")
    finally:


        if _tempos:
            registrar_log([f"⏱️  {pedido}: " + " | ".join(_tempos)])

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# CICLO DE PROCESSAMENTO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def processar_ciclo(retomar=False):
    _recarregar_xpaths()

    if not os.path.exists(CAMINHO_FILA):
        registrar_log(["⛔  Fila_Aprovacao.xlsx não encontrada."])
        return

    try:
        wb_in = load_workbook(CAMINHO_FILA, data_only=True)
        ws_in = wb_in[ABA_ENTRADA] if ABA_ENTRADA in wb_in.sheetnames else wb_in.active
    except Exception as e:
        registrar_log([f"❌  Erro ao ler Fila_Aprovacao.xlsx: {e}"])
        registrar_erro_permanente(f"LER FILA -- {e}")
        return

    cabecalho = [str(c.value).strip() if c.value is not None else "" for c in ws_in[1]]
    col_map = {nome: idx + 1 for idx, nome in enumerate(cabecalho) if nome}


    col_pedido = col_map.get("PEDIDO")
    ultima_linha_real = 1
    if col_pedido:
        for linha_num in range(2, ws_in.max_row + 1):
            valor = ws_in.cell(row=linha_num, column=col_pedido).value
            if valor is not None and str(valor).strip():
                ultima_linha_real = linha_num
    else:
        ultima_linha_real = ws_in.max_row

    total = max(0, ultima_linha_real - 1)
    registrar_log([f"📊  Iniciando processamento de {total} pedido(s)."])


    resultado_anterior = {}
    if retomar and os.path.exists(CAMINHO_RESULTADO):
        try:
            wb_ant = load_workbook(CAMINHO_RESULTADO, data_only=True)
            ws_ant = wb_ant.active
            cabecalho_ant = {idx: str(c.value).strip() if c.value is not None else ""
                             for idx, c in enumerate(ws_ant[1], start=1)}
            col_pedido_ant = next((i for i, n in cabecalho_ant.items() if n == "PEDIDO"), None)
            col_status_ant = next((i for i, n in cabecalho_ant.items() if n == "STATUS"), None)
            if col_pedido_ant and col_status_ant:
                for r in range(2, ws_ant.max_row + 1):
                    ped_str = str(ws_ant.cell(row=r, column=col_pedido_ant).value or "").strip()
                    status_val = ws_ant.cell(row=r, column=col_status_ant).value
                    if ped_str and status_val and str(status_val).strip():
                        resultado_anterior[ped_str] = {
                            nome: ws_ant.cell(row=r, column=idx).value
                            for idx, nome in cabecalho_ant.items() if nome
                        }
            if resultado_anterior:
                registrar_log([f"🔁  Retomando ciclo -- {len(resultado_anterior)} pedido(s) já resolvido(s) "
                               f"numa tentativa anterior não serão reprocessados."])
        except Exception as e:
            registrar_erro_permanente(f"RETOMAR CICLO -- {type(e).__name__}: {e}")
            resultado_anterior = {}

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = ABA_ENTRADA
    for nome, idx in col_map.items():
        ws_out.cell(row=1, column=idx, value=nome)

    driver = iniciar_navegador()
    if not driver:
        _gravar_status_ciclo("ATENCAO", "Falha ao abrir o navegador.")
        return

    ok_login, erro_login = fazer_login(driver)
    if not ok_login:
        registrar_log([f"❌  Falha no login: {erro_login}"])
        registrar_erro_permanente(f"LOGIN -- {erro_login}")
        _gravar_status_ciclo("ATENCAO", f"Falha no login: {erro_login}")
        try: driver.quit()
        except Exception: pass
        return

    if not navegar_para_pedidos(driver):
        _gravar_status_ciclo("ATENCAO", "Falha ao navegar até a tela de Pedidos.")
        try: driver.quit()
        except Exception: pass
        return

    linhas_aprovadas = []
    parado_criticamente = False
    navegador_morreu = False
    for linha_num in range(2, ultima_linha_real + 1):
        dados = {nome: ws_in.cell(row=linha_num, column=idx).value for nome, idx in col_map.items()}
        for nome, idx in col_map.items():
            ws_out.cell(row=linha_num, column=idx, value=dados.get(nome))

        ped_str = str(dados.get("PEDIDO", "") or "").strip()
        if ped_str and ped_str in resultado_anterior:
            linha_anterior = resultado_anterior[ped_str]
            for nome, idx in col_map.items():
                if nome in linha_anterior:
                    ws_out.cell(row=linha_num, column=idx, value=linha_anterior[nome])
            status_ant = str(linha_anterior.get("STATUS", "") or "")
            registrar_log([f"⏭️  {ped_str}: já resolvido numa tentativa anterior (status: {status_ant}) -- pulando."])
            if status_ant.startswith("APROVADO"):
                linhas_aprovadas.append(linha_num)
            continue


        try:
            driver.title
        except WebDriverException as e:
            registrar_log([
                f"🛑  CICLO INTERROMPIDO: a sessão do navegador morreu ({type(e).__name__}) "
                f"antes de processar a linha {linha_num}.",
                "Os pedidos restantes desta fila NÃO foram processados -- "
                "o robô vai tentar de novo automaticamente, com um navegador novo.",
            ])
            registrar_erro_permanente(f"CICLO INTERROMPIDO -- sessão do navegador morreu ({type(e).__name__}: {e})")
            parado_criticamente = True
            navegador_morreu = True
            break

        try:
            processar_pedido(driver, dados, ws_out, linha_num, col_map)
        except ParadaCritica as e:
            registrar_log([
                f"🛑  CICLO INTERROMPIDO: {e}",
                "Os pedidos restantes desta fila NÃO foram processados -- "
                "o robô vai tentar de novo no próximo gatilho.",
            ])
            registrar_erro_permanente(f"CICLO INTERROMPIDO -- {e}")
            parado_criticamente = True
            break
        except WebDriverException as e:
            registrar_log([
                f"🛑  CICLO INTERROMPIDO: a sessão do navegador morreu ({type(e).__name__}) "
                f"processando a linha {linha_num}.",
                "Os pedidos restantes desta fila NÃO foram processados -- "
                "o robô vai tentar de novo automaticamente, com um navegador novo.",
            ])
            registrar_erro_permanente(f"CICLO INTERROMPIDO -- sessão do navegador morreu ({type(e).__name__}: {e})")
            parado_criticamente = True
            navegador_morreu = True
            break
        except Exception as e:
            registrar_erro_permanente(f"PROCESSAR LINHA {linha_num} -- {type(e).__name__}: {e}")

        status_gravado = ws_out.cell(row=linha_num, column=col_map.get("STATUS", 1)).value if "STATUS" in col_map else None
        if status_gravado and str(status_gravado).startswith("APROVADO"):
            linhas_aprovadas.append(linha_num)

    try:
        driver.quit()
    except Exception:
        pass

    try:
        os.makedirs(os.path.dirname(CAMINHO_RESULTADO), exist_ok=True)
        wb_out.save(CAMINHO_RESULTADO)
        if parado_criticamente:
            registrar_log(["⚠️  Ciclo salvo PARCIALMENTE (interrompido antes do fim).",
                          f"Resultado salvo em: {CAMINHO_RESULTADO}"])


            if not navegador_morreu:
                _gravar_status_ciclo("ATENCAO", "Ciclo interrompido antes do fim -- nem todos os pedidos da fila foram processados.")
        else:
            registrar_log(["🏁  Ciclo concluído.", f"Resultado salvo em: {CAMINHO_RESULTADO}"])
            _gravar_status_ciclo("OK", "Ciclo concluído normalmente.")
    except Exception as e:
        registrar_log([f"❌  Erro ao salvar Resultado_Aprovacao.xlsx: {e}"])
        registrar_erro_permanente(f"SALVAR RESULTADO -- {e}")
        _gravar_status_ciclo("ATENCAO", f"Erro ao salvar o resultado: {e}")

    if linhas_aprovadas:
        try:
            _atualizar_historico(ws_out, col_map, linhas_aprovadas)
        except Exception as e:
            registrar_erro_permanente(f"HISTORICO -- {e}")


    if total > 0:
        _salvar_historico_impressao(wb_in, wb_out)


    return navegador_morreu


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# JANELA DE LANÇAMENTO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
class JanelaLancamento:
    def __init__(self):
        self.aprovado      = False
        self._falha_lib    = False
        self._falha_chrome = False

        self.root = ctk.CTk()
        self.root.title("Robô Aprovador")
        self.root.configure(fg_color=COR_FUNDO)
        self.root.resizable(False, False)
        self.root.attributes("-topmost", True)
        self.root.protocol("WM_DELETE_WINDOW", self._fechar)

        self._construir_ui()
        self._centralizar(560, 620)
        threading.Thread(target=self._verificar_lock_inicial, daemon=True).start()
        self.root.mainloop()

    def _centralizar(self, w, h):
        self.root.update_idletasks()
        sw = self.root.winfo_screenwidth()
        sh = self.root.winfo_screenheight()
        self.root.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")

    def _construir_ui(self):
        hdr = ctk.CTkFrame(self.root, fg_color="transparent")
        hdr.pack(fill="x", padx=24, pady=(24, 16))

        badge = ctk.CTkLabel(hdr, text="✅", font=("Segoe UI", 20),
                             fg_color=COR_HEADER, text_color=COR_HEADER_TXT,
                             corner_radius=12, width=44, height=44)
        badge.pack(side="left")

        txt_hdr = ctk.CTkFrame(hdr, fg_color="transparent")
        txt_hdr.pack(side="left", padx=(12, 0))
        ctk.CTkLabel(txt_hdr, text="Robô Aprovador", font=("Segoe UI", 17, "bold"),
                    text_color=COR_TEXTO, anchor="w").pack(anchor="w")
        ctk.CTkLabel(txt_hdr, text="Aprovação automática de pedidos", font=("Segoe UI", 12),
                    text_color=COR_DETALHE, anchor="w").pack(anchor="w")

        ctk.CTkButton(
            hdr, text=f"👤  CRIADO POR: {NOME_CRIADOR}", font=("Segoe UI", 9),
            fg_color="transparent", text_color=COR_MUTED, hover_color=COR_CARD,
            corner_radius=6, height=22, anchor="e",
            command=lambda: webbrowser.open(LINK_GITHUB_CRIADOR)
        ).pack(side="right", anchor="n")

        self.frm_controle = ctk.CTkFrame(self.root, fg_color="transparent")
        self.frm_controle.pack(fill="both", expand=True, padx=24)

        self.frm_btn = ctk.CTkFrame(self.root, fg_color="transparent")
        self.frm_btn.pack(fill="x", padx=24, pady=(12, 20))
        ctk.CTkButton(self.frm_btn, text="✖  Fechar", font=("Segoe UI", 11),
                     fg_color=COR_BTN_C, text_color=COR_BTN_C_TXT, hover_color=COR_BORDA,
                     corner_radius=8, height=36,
                     command=self._fechar).pack(side="right")

    def _limpar_controle(self):
        for w in self.frm_controle.winfo_children():
            w.destroy()

    def _card(self, titulo=None):
        card = ctk.CTkFrame(self.frm_controle, fg_color=COR_CARD, corner_radius=12,
                            border_width=1, border_color=COR_BORDA)
        card.pack(fill="x", pady=(0, 12))
        interno = ctk.CTkFrame(card, fg_color="transparent")
        interno.pack(fill="x", padx=16, pady=14)
        if titulo:
            ctk.CTkLabel(interno, text=titulo.upper(), font=("Segoe UI", 10, "bold"),
                        text_color=COR_MUTED, anchor="w").pack(anchor="w", pady=(0, 8))
        return interno

    def _ir_para_etapa(self, etapa):
        self._limpar_controle()
        getattr(self, f"_etapa_{etapa}")()

    def _mostrar_continuar(self, spoiler_texto, proxima_etapa, texto_botao="Continuar  →"):
        frm = ctk.CTkFrame(self.frm_controle, fg_color="transparent")
        frm.pack(fill="x", pady=(14, 0))
        if spoiler_texto:
            ctk.CTkLabel(frm, text=spoiler_texto, font=("Segoe UI", 10),
                        text_color=COR_MUTED, anchor="w").pack(anchor="w", pady=(0, 6))
        ctk.CTkButton(frm, text=texto_botao, font=("Segoe UI", 12, "bold"),
                     fg_color=COR_HEADER, text_color=COR_HEADER_TXT, hover_color="#E6B325",
                     corner_radius=8, height=38,
                     command=lambda: self._ir_para_etapa(proxima_etapa)).pack(fill="x")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 0 -- início
    # ─── ⋆⋅☆⋅⋆ ───
    def _verificar_lock_inicial(self):
        time.sleep(0.3)
        pode_rodar, info_lock = _verificar_instancia_unica()
        if not pode_rodar:
            self.root.after(0, lambda: self._tela_bloqueado(info_lock))
            return
        self.root.after(0, lambda: self._ir_para_etapa("bibliotecas"))

    def _tela_bloqueado(self, info_lock):
        self._limpar_controle()
        card = self._card()
        ctk.CTkLabel(card, text="⚠️  ROBÔ JÁ ATIVO EM OUTRA MÁQUINA", font=("Segoe UI", 12, "bold"),
                    text_color=COR_AVISO, anchor="w").pack(anchor="w")
        ctk.CTkLabel(card, text=f"Em execução em: {info_lock}", font=("Segoe UI", 10),
                    text_color=COR_MUTED, anchor="w", wraplength=460,
                    justify="left").pack(anchor="w", pady=(6, 0))
        ctk.CTkLabel(card, text="Feche esta janela.", font=("Segoe UI", 10),
                    text_color=COR_MUTED, anchor="w").pack(anchor="w")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 1 -- bibliotecas
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_bibliotecas(self):
        card = self._card("Bibliotecas")
        self._frm_bibliotecas = ctk.CTkFrame(card, fg_color="transparent")
        self._frm_bibliotecas.pack(fill="x")

        self._lbl_status_geral = ctk.CTkLabel(card, text="Verificando...", font=("Segoe UI", 10),
                                              text_color=COR_MUTED, anchor="w")
        self._lbl_status_geral.pack(fill="x", pady=(10, 4))

        self._barra_progresso = ctk.CTkProgressBar(card, height=6, corner_radius=4,
                                                    fg_color=COR_BORDA, progress_color=COR_HEADER)
        self._barra_progresso.set(0)
        self._barra_progresso.pack(fill="x")
        threading.Thread(target=self._checar_bibliotecas, daemon=True).start()

    def _status_geral(self, texto):
        try:
            self.root.after(0, lambda: self._lbl_status_geral.configure(text=texto))
        except Exception:
            pass

    def _fade_in_label(self, widget, cor_final, passos=10, intervalo=18):
        def _hex_rgb(c):
            c = c.lstrip("#")
            return tuple(int(c[i:i+2], 16) for i in (0, 2, 4))
        r1, g1, b1 = _hex_rgb(COR_CARD)
        r2, g2, b2 = _hex_rgb(cor_final)
        def _tick(i=0):
            if i > passos:
                try:
                    widget.configure(text_color=cor_final)
                except Exception:
                    pass
                return
            frac = i / passos
            cor = "#{:02X}{:02X}{:02X}".format(
                int(r1 + (r2 - r1) * frac), int(g1 + (g2 - g1) * frac), int(b1 + (b2 - b1) * frac))
            try:
                widget.configure(text_color=cor)
            except Exception:
                return
            self.root.after(intervalo, lambda: _tick(i + 1))
        _tick()

    def _linha_biblioteca(self, nome, estado="instalando"):
        resultado = {}
        evento = threading.Event()
        def _run():
            try:
                linha = ctk.CTkFrame(self._frm_bibliotecas, fg_color="transparent")
                linha.pack(fill="x")
                interno = ctk.CTkFrame(linha, fg_color="transparent")
                interno.pack(fill="x", pady=8)

                canvas_icone = tk.Canvas(interno, width=16, height=16, bg=COR_CARD,
                                         highlightthickness=0)
                canvas_icone.pack(side="left", padx=(0, 8))

                lbl_nome = ctk.CTkLabel(interno, text=nome, font=("Segoe UI", 12),
                                        text_color=COR_CARD)
                lbl_nome.pack(side="left")

                lbl_resultado = ctk.CTkLabel(interno, text="", font=("Segoe UI", 13, "bold"),
                                             text_color=COR_CARD)
                lbl_resultado.pack(side="right")

                ctk.CTkFrame(self._frm_bibliotecas, fg_color=COR_BORDA, height=1).pack(fill="x")

                resultado["canvas_icone"] = canvas_icone
                resultado["lbl_nome"] = lbl_nome
                resultado["lbl_resultado"] = lbl_resultado
                resultado["spinner_ativo"] = {"ativo": False}

                self._fade_in_label(lbl_nome, COR_TEXTO)
            except Exception:
                pass
            evento.set()
        try:
            self.root.after(0, _run)
        except Exception:
            evento.set()
        evento.wait(timeout=2)
        if resultado:
            self._definir_estado_linha(resultado, estado)
        return resultado

    def _girar_spinner(self, canvas_spinner, ativo_ref):
        angulo = [0]
        def _tick():
            if not ativo_ref["ativo"] or not canvas_spinner:
                return
            angulo[0] = (angulo[0] + 24) % 360
            try:
                canvas_spinner.itemconfig("arco", start=angulo[0])
            except Exception:
                return
            self.root.after(40, _tick)
        try:
            self.root.after(0, _tick)
        except Exception:
            pass

    def _definir_estado_linha(self, referencia, estado):
        def _run():
            try:
                canvas = referencia["canvas_icone"]
                referencia["spinner_ativo"]["ativo"] = False
                canvas.delete("all")

                if estado == "instalando":
                    canvas.create_arc(2, 2, 14, 14, start=0, extent=270, style="arc",
                                      outline=COR_HEADER, width=2, tags="arco")
                    referencia["spinner_ativo"]["ativo"] = True
                    self._girar_spinner(canvas, referencia["spinner_ativo"])
                else:
                    canvas.create_oval(2, 2, 14, 14, outline=COR_HEADER, width=2)

                if estado in ("ok", "instalado"):
                    referencia["lbl_resultado"].configure(text="✓", text_color=COR_OK)
                elif estado == "erro":
                    referencia["lbl_resultado"].configure(text="✕", text_color=COR_ERRO)
            except Exception:
                pass
        try:
            self.root.after(0, _run)
        except Exception:
            pass

    def _checar_bibliotecas(self):
        total = len(_STATUS_INSTALL)
        prontas = 0

        def _atualizar_progresso():
            try:
                self.root.after(0, lambda: self._barra_progresso.set(prontas / total if total else 1))
            except Exception:
                pass

        pendentes = []
        for nome_pip, status in _STATUS_INSTALL.items():
            if status == "ok":
                self._status_geral(f"Verificando {nome_pip}...")
                self._linha_biblioteca(nome_pip, "ok")
                prontas += 1
                _atualizar_progresso()
                time.sleep(0.35)
            else:
                pendentes.append(nome_pip)

        for nome_pip in pendentes:
            modulo = next((ni for ni, np in _LIBS if np == nome_pip), nome_pip)
            self._status_geral(f"Instalando {nome_pip}...")
            referencia = self._linha_biblioteca(nome_pip, "instalando")
            try:
                subprocess.run([sys.executable, "-m", "pip", "install", nome_pip, "--quiet"],
                               capture_output=True, timeout=180, creationflags=_NO_WINDOW)
                importlib.invalidate_caches()
                __import__(modulo)
                _STATUS_INSTALL[nome_pip] = "ok"
                self._definir_estado_linha(referencia, "instalado")
                self._status_geral(f"{nome_pip} instalado!")
            except Exception as e_lib:
                _STATUS_INSTALL[nome_pip] = "erro"
                self._definir_estado_linha(referencia, "erro")
                self._status_geral(f"Falha ao instalar {nome_pip}")
                registrar_erro_permanente(f"INSTALAR {nome_pip} -- {e_lib}")
            prontas += 1
            _atualizar_progresso()
            time.sleep(0.2)

        if not _IMPORTS_OK:
            _tentar_importar_reais()

        self._falha_lib = not _IMPORTS_OK
        self._status_geral("Tudo pronto!" if _IMPORTS_OK else "Há um problema a resolver.")
        self.root.after(500, self._mostrar_proximo_bibliotecas)

    def _mostrar_proximo_bibliotecas(self):
        if _IMPORTS_OK:
            self._mostrar_continuar("Próximo: Pasta Bin (Chrome/Driver)", "bin")
            return
        frm = ctk.CTkFrame(self.frm_controle, fg_color="transparent")
        frm.pack(fill="x", pady=(14, 0))
        ctk.CTkLabel(frm, text=f"❌  Erro: {_IMPORTS_ERRO[:150]}", font=("Segoe UI", 10),
                    text_color=COR_ERRO, wraplength=460, justify="left").pack(anchor="w", pady=(0, 8))
        ctk.CTkLabel(frm, text="Algumas bibliotecas só funcionam de verdade depois "
                              "de reiniciar o programa.",
                    font=("Segoe UI", 9), text_color=COR_MUTED,
                    wraplength=460, justify="left").pack(anchor="w", pady=(0, 8))
        ctk.CTkButton(frm, text="🔄  Reiniciar", font=("Segoe UI", 11, "bold"),
                     fg_color=COR_BTN_C, text_color=COR_BTN_C_TXT, hover_color=COR_BORDA,
                     corner_radius=8, height=36,
                     command=self._reiniciar_programa).pack(fill="x")

    def _reiniciar_programa(self):
        try:
            subprocess.Popen([sys.executable, os.path.abspath(__file__)])
        except Exception as e:
            registrar_erro_permanente(f"REINICIAR -- {e}")
        finally:
            _liberar_lock()
            try:
                self.root.destroy()
            except Exception:
                pass
            os._exit(0)


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 2 -- bin (Chrome/Driver)
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_bin(self):
        card = self._card("Pasta Bin (Chrome/Driver locais)")
        self._lbl_bin_status = ctk.CTkLabel(card, text="Verificando...", font=("Segoe UI", 11),
                                            text_color=COR_TEXTO, anchor="w")
        self._lbl_bin_status.pack(anchor="w")
        self._lbl_bin_detalhe = ctk.CTkLabel(card, text="", font=("Segoe UI", 9),
                                             text_color=COR_MUTED, anchor="w",
                                             wraplength=460, justify="left")
        self._lbl_bin_detalhe.pack(anchor="w", pady=(4, 0))
        threading.Thread(target=self._checar_bin, daemon=True).start()

    def _checar_bin(self):
        if _bin_completa():
            self.root.after(0, lambda: self._lbl_bin_status.configure(
                text="✅  Já existe localmente.", text_color=COR_OK))
            self.root.after(300, lambda: self._mostrar_continuar("Próximo: Chrome e Driver", "chrome_driver"))
            return

        self.root.after(0, lambda: self._lbl_bin_status.configure(
            text="⏳  Copiando da rede...", text_color=COR_AVISO))
        self.root.after(0, lambda: self._lbl_bin_detalhe.configure(
            text=f"Origem: {CAMINHO_REDE_BIN}\nIsso pode demorar vários minutos, aguarde..."))

        parar_ticker = threading.Event()
        inicio_copia = time.time()

        def _ticker():
            while not parar_ticker.is_set():
                decorrido = int(time.time() - inicio_copia)
                try:
                    self.root.after(0, lambda d=decorrido: self._lbl_bin_status.configure(
                        text=f"⏳  Copiando Bin... {d//60}min {d%60}s"))
                except Exception:
                    return
                time.sleep(1)

        threading.Thread(target=_ticker, daemon=True).start()
        ok_copia, erro_copia = _copiar_bin_da_rede()
        parar_ticker.set()

        if ok_copia:
            self.root.after(0, lambda: self._lbl_bin_status.configure(
                text="✅  Bin copiada com sucesso.", text_color=COR_OK))
            self.root.after(0, lambda: self._lbl_bin_detalhe.configure(text=""))
        else:
            self.root.after(0, lambda: self._lbl_bin_status.configure(
                text="❌  Falha ao copiar Bin da rede.", text_color=COR_ERRO))
            self.root.after(0, lambda: self._lbl_bin_detalhe.configure(text=erro_copia[:200]))
            registrar_erro_permanente(f"CÓPIA BIN -- {erro_copia}")

        self.root.after(300, lambda: self._mostrar_continuar("Próximo: Chrome e Driver", "chrome_driver"))


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 3 -- Chrome e Driver
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_chrome_driver(self):
        card = self._card("Chrome e Driver")
        cr = _checar_chrome()

        linhas = []
        linhas.append(("✅  Chrome encontrado (local)", COR_OK) if cr["chrome_existe"]
                      else ("❌  Chrome não encontrado", COR_ERRO))
        linhas.append(("✅  Driver encontrado (local)", COR_OK) if cr["driver_existe"]
                      else ("❌  Driver não encontrado", COR_ERRO))
        if cr["compat"]:
            linhas.append(("✅  Compatíveis", COR_OK))

        for texto, cor in linhas:
            ctk.CTkLabel(card, text=texto, font=("Segoe UI", 11), text_color=cor,
                        anchor="w").pack(anchor="w", pady=2)

        if not cr["chrome_existe"] or not cr["driver_existe"]:
            for l in cr["erro"].split("\n"):
                if l.strip():
                    ctk.CTkLabel(card, text=l.strip(), font=("Segoe UI", 9),
                                text_color=COR_MUTED, anchor="w").pack(anchor="w")

        self._falha_chrome = not cr["compat"]
        self._mostrar_continuar("Pronto para iniciar!", "final")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA FINAL -- iniciar
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_final(self):
        card = self._card()
        falha = self._falha_lib or self._falha_chrome
        if falha:
            ctk.CTkLabel(card, text="⚠️  Alguns itens tiveram problema -- revise antes de iniciar.",
                        font=("Segoe UI", 11), text_color=COR_AVISO, anchor="w",
                        wraplength=460, justify="left").pack(anchor="w")
        else:
            ctk.CTkLabel(card, text="✅  Tudo certo! Pronto para iniciar.",
                        font=("Segoe UI", 13, "bold"), text_color=COR_OK, anchor="w").pack(anchor="w")

        if not falha:
            ctk.CTkButton(self.frm_controle, text="▶  Iniciar", font=("Segoe UI", 12, "bold"),
                         fg_color=COR_HEADER, text_color=COR_HEADER_TXT, hover_color="#E6B325",
                         corner_radius=8, height=38,
                         command=self._iniciar).pack(fill="x", pady=(6, 0))

    def _iniciar(self):
        self.aprovado = True
        self.root.destroy()

    def _fechar(self):
        self.aprovado = False
        _liberar_lock()
        self.root.destroy()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LOOP PRINCIPAL
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _abrir_log(icon=None, item=None):
    try:
        os.startfile(LOG_ARQUIVO)
    except Exception as e:
        registrar_erro_permanente(f"ABRIR LOG -- {e}")


def _criar_icone_tray():
    try:
        from PIL import Image, ImageDraw
        img  = Image.new("RGBA", (64, 64), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        draw.ellipse([0, 0, 63, 63],   fill="#18181B")
        draw.ellipse([16, 16, 47, 47], fill="#FFC72C")
        draw.ellipse([24, 24, 39, 39], fill="#18181B")
        return img
    except Exception:
        from PIL import Image
        return Image.new("RGBA", (64, 64), "#18181B")

def _esta_em_pausa(agora=None):
    agora = agora or datetime.now()
    hora_atual = agora.time()
    if HORA_ENCERRAR <= HORA_RETOMAR:
        return HORA_ENCERRAR <= hora_atual < HORA_RETOMAR
    return hora_atual >= HORA_ENCERRAR or hora_atual < HORA_RETOMAR

DIAS_MANTER_DIAGNOSTICO       = 10
INTERVALO_LIMPEZA_DIAGNOSTICO = 24 * 60 * 60

def _limpar_arquivos_diagnostico_antigos():
    try:
        pasta_erro_base = os.path.join(_pasta_logs, "ERRO")
        if not os.path.isdir(pasta_erro_base):
            return
        limite = time.time() - (DIAS_MANTER_DIAGNOSTICO * 24 * 60 * 60)
        removidos = 0
        for ano in os.listdir(pasta_erro_base):
            pasta_ano = os.path.join(pasta_erro_base, ano)
            if not os.path.isdir(pasta_ano):
                continue
            for mes in os.listdir(pasta_ano):
                pasta_mes = os.path.join(pasta_ano, mes)
                if not os.path.isdir(pasta_mes):
                    continue
                for dia in os.listdir(pasta_mes):
                    pasta_dia = os.path.join(pasta_mes, dia)
                    if not os.path.isdir(pasta_dia):
                        continue
                    try:
                        if os.path.getmtime(pasta_dia) < limite:
                            shutil.rmtree(pasta_dia)
                            removidos += 1
                    except Exception:
                        continue
                try:
                    if not os.listdir(pasta_mes):
                        os.rmdir(pasta_mes)
                except Exception:
                    pass
            try:
                if not os.listdir(pasta_ano):
                    os.rmdir(pasta_ano)
            except Exception:
                pass
        if removidos:
            registrar_log([f"🧹  Limpeza automática: {removidos} pasta(s) de dia com mais "
                           f"de {DIAS_MANTER_DIAGNOSTICO} dias removidas."])
    except Exception as e:
        registrar_erro_permanente(f"LIMPEZA DIAGNÓSTICO -- {e}")

def _loop_aprovacao_interno():
    global _status_atual

    _gerenciar_log()
    registrar_log([
        "🤖  ROBÔ APROVADOR -- SISTEMA INICIADO",
        f"📅  Data: {_agora_completo()}",
        f"👀  Vigiando: {CAMINHO_FILA}",
        f"PID: {os.getpid()} | PC: {os.environ.get('COMPUTERNAME','?')}",
        f"Pausa noturna: {HORA_ENCERRAR.strftime('%H:%M')} -- {HORA_RETOMAR.strftime('%H:%M')} (retomada automática)",
    ])
    _status_atual = "Ativo"

    ultimo_mtime = os.path.getmtime(CAMINHO_FILA) if os.path.exists(CAMINHO_FILA) else 0
    ultimo_heartbeat = time.time()
    ultima_limpeza   = 0.0
    estava_pausado = False

    _limpar_arquivos_diagnostico_antigos()
    ultima_limpeza = time.time()

    while not _parar_loop.is_set():
        agora    = datetime.now()
        em_pausa = _esta_em_pausa(agora)

        if time.time() - ultima_limpeza >= INTERVALO_LIMPEZA_DIAGNOSTICO:
            _limpar_arquivos_diagnostico_antigos()
            ultima_limpeza = time.time()

        if em_pausa:
            if not estava_pausado:
                registrar_log([
                    f"PAUSA -- Encerramento automático às {HORA_ENCERRAR.strftime('%H:%M')}h.",
                    f"Data: {_agora_completo()}",
                    f"Retoma automaticamente às {HORA_RETOMAR.strftime('%H:%M')}h.",
                ])
                estava_pausado = True
            _status_atual = f"🌙 Pausado até {HORA_RETOMAR.strftime('%H:%M')}"

            if time.time() - ultimo_heartbeat >= INTERVALO_HEARTBEAT:
                _atualizar_heartbeat_lock()
                ultimo_heartbeat = time.time()

            for _ in range(30):
                if _parar_loop.is_set(): break
                time.sleep(1)
            continue

        if estava_pausado:
            registrar_log([
                f"RETOMANDO -- Operação normal retomada às {HORA_RETOMAR.strftime('%H:%M')}h.",
                f"Data: {_agora_completo()}",
            ])
            estava_pausado = False
            _status_atual = "Ativo"

        try:
            if os.path.exists(CAMINHO_FILA):
                atual_mtime = os.path.getmtime(CAMINHO_FILA)
                if atual_mtime != ultimo_mtime:
                    registrar_log(["⚡  Gatilho detectado -- Fila_Aprovacao.xlsx atualizada.",
                                   f"Data: {_agora_completo()}"])
                    _status_atual = "Processando..."
                    time.sleep(3)
                    navegador_morreu = processar_ciclo()


                    tentativas_navegador = 0
                    while (navegador_morreu and tentativas_navegador < LIMITE_RETRY_NAVEGADOR_MORTO
                           and not _parar_loop.is_set()):
                        tentativas_navegador += 1
                        registrar_log([
                            f"🔁  Tentando de novo automaticamente (navegador morreu) -- "
                            f"tentativa {tentativas_navegador}/{LIMITE_RETRY_NAVEGADOR_MORTO}...",
                        ])
                        time.sleep(10)
                        navegador_morreu = processar_ciclo(retomar=True)
                    if navegador_morreu:
                        registrar_log([
                            f"⚠️  Navegador continuou morrendo depois de {LIMITE_RETRY_NAVEGADOR_MORTO} "
                            f"tentativas automáticas -- novas tentativas suspensas. Verificar manualmente "
                            f"(memória/antivírus/versão do Chrome) e alterar a planilha novamente para uma nova tentativa.",
                        ])
                        registrar_erro_permanente(
                            f"NAVEGADOR MORREU {LIMITE_RETRY_NAVEGADOR_MORTO + 1}x SEGUIDAS -- tentativas suspensas."
                        )


                        _gravar_status_ciclo(
                            "ATENCAO",
                            f"Navegador morreu {LIMITE_RETRY_NAVEGADOR_MORTO + 1}x seguidas -- ciclo não foi concluído."
                        )


                    ultimo_mtime = os.path.getmtime(CAMINHO_FILA) if os.path.exists(CAMINHO_FILA) else atual_mtime
                    _status_atual = f"OK -- {datetime.now().strftime('%H:%M')}"
        except Exception as e:
            _status_atual = f"Erro -- {datetime.now().strftime('%H:%M')}"
            registrar_log(["💥  Erro crítico no loop principal", f"Detalhe: {e}"])
            registrar_erro_permanente(f"LOOP -- {e}")

        if time.time() - ultimo_heartbeat >= INTERVALO_HEARTBEAT:
            _atualizar_heartbeat_lock()
            ultimo_heartbeat = time.time()

        for _ in range(5):
            if _parar_loop.is_set(): break
            time.sleep(1)

def _loop_aprovacao():
    while not _parar_loop.is_set():
        try:
            _loop_aprovacao_interno()
            break
        except Exception as e:
            try:
                import traceback
                registrar_erro_permanente(f"LOOP MORREU SEM TRATAMENTO -- reiniciando automaticamente. {e}\n{traceback.format_exc()}")
                registrar_log([
                    "💥  O loop principal parou por um erro inesperado -- reiniciando automaticamente",
                    f"🕐  {_agora_completo()}",
                    f"     Detalhe: {e}",
                ])
            except Exception:
                pass
            for _ in range(30):
                if _parar_loop.is_set(): break
                time.sleep(1)
    _liberar_lock()

def _iniciar_tray():
    try:
        import pystray
    except ImportError:
        _loop_aprovacao()
        return

    def texto_status(item):
        return f"Status: {_status_atual}"

    def parar(icon, item):
        _parar_loop.set()
        _liberar_lock()
        registrar_log(["Robô parado pelo usuário.", f"Data: {_agora_completo()}"])
        threading.Timer(5, os._exit, args=(0,)).start()
        icon.stop()

    menu = pystray.Menu(
        pystray.MenuItem("📄  Abrir log", _abrir_log, default=True),
        pystray.MenuItem(texto_status, None, enabled=False),
        pystray.Menu.SEPARATOR,
        pystray.MenuItem("⏹  Parar robô", parar),
    )
    icone = pystray.Icon("RoboAprovador", _criar_icone_tray(), "Robô Aprovador", menu=menu)

    threading.Thread(target=_loop_aprovacao, daemon=True).start()
    icone.run()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# PONTO DE ENTRADA
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
if __name__ == "__main__":
    try:
        janela = JanelaLancamento()
    except Exception as _e:
        import traceback
        _erro_completo = traceback.format_exc()
        try:
            with open(os.path.join(pasta_script, "erro_fatal.txt"), "w", encoding="utf-8") as _f:
                _f.write(f"[{datetime.now().strftime('%d/%m/%Y %H:%M:%S')}]\n{_erro_completo}")
        except Exception:
            pass
        try:
            import tkinter.messagebox as _mb
            _raiz_erro = tk.Tk()
            _raiz_erro.withdraw()
            _mb.showerror(
                "Robô Aprovador -- Erro ao iniciar",
                f"O robô encontrou um erro e não conseguiu abrir:\n\n{_e}\n\n"
                f"Detalhes salvos em: erro_fatal.txt (nesta mesma pasta)."
            )
            _raiz_erro.destroy()
        except Exception:
            pass
        sys.exit(1)

    if not janela.aprovado:
        sys.exit(0)
    _iniciar_tray()
