# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# BIBLIOTECAS
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
import subprocess
import sys
import os
import time
import glob
import shutil
import threading
import importlib
import webbrowser
from datetime import datetime
import tkinter as tk

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

pasta_script    = os.path.dirname(os.path.abspath(__file__))
pasta_downloads = os.path.join(os.path.expanduser("~"), "Downloads")
_NO_WINDOW      = getattr(subprocess, "CREATE_NO_WINDOW", 0)


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# DETECÇÃO PARA AUTO-INSTALAÇÃO DE BIBLIOTECAS
# Instalação ao vivo em _checar_bibliotecas (tela Bibliotecas).
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_LIBS = [
    ("pandas",        "pandas"),
    ("openpyxl",      "openpyxl"),
    ("win32com",      "pywin32"),
    ("dotenv",        "python-dotenv"),
    ("selenium",      "selenium"),
    ("pystray",       "pystray"),
    ("PIL",           "Pillow"),
    ("customtkinter", "customtkinter"),
]
_STATUS_INSTALL = {}
for _ni, _np in _LIBS:
    try:
        __import__(_ni)
        _STATUS_INSTALL[_np] = "ok"
    except ImportError:
        _STATUS_INSTALL[_np] = "faltando"


if _STATUS_INSTALL.get("customtkinter") == "faltando":
    _splash = tk.Tk()
    _splash.title("Base Geral")
    _splash.configure(bg="#18181B")
    _splash.resizable(False, False)
    _sw, _sh = _splash.winfo_screenwidth(), _splash.winfo_screenheight()
    _splash.geometry(f"360x110+{(_sw-360)//2}+{(_sh-110)//2}")
    tk.Label(_splash, text="Preparando o ambiente pela primeira vez...",
             font=("Segoe UI", 11), fg="#F2F2F2", bg="#18181B").pack(pady=(28, 6))
    tk.Label(_splash, text="Instalando customtkinter...",
             font=("Segoe UI", 9), fg="#A8A8A8", bg="#18181B").pack()
    _splash.update()
    try:
        subprocess.run([sys.executable, "-m", "pip", "install", "customtkinter", "--quiet"],
                       capture_output=True, timeout=180, creationflags=_NO_WINDOW)
        __import__("customtkinter")
        _STATUS_INSTALL["customtkinter"] = "ok"
    except Exception:
        _STATUS_INSTALL["customtkinter"] = "faltando"
    _splash.destroy()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# IMPORTS EXTERNOS
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_IMPORTS_OK = True
_IMPORTS_ERRO = ""

# DLLs do pywin32 no path > sem isso, o win32com falha depois de instalar o pywin32 em tempo de execução.
def _corrigir_dll_pywin32():
    try:
        import site
        pastas = list(site.getsitepackages())
        try:
            pastas.append(site.getusersitepackages())
        except Exception:
            pass
        for base in pastas:
            pasta_dll = os.path.join(base, "pywin32_system32")
            if os.path.isdir(pasta_dll):
                try:
                    os.add_dll_directory(pasta_dll)
                except Exception:
                    pass
    except Exception:
        pass

def _tentar_importar_reais():
    global _IMPORTS_OK, _IMPORTS_ERRO
    global ctk, pd, win32, load_dotenv, webdriver, Options, Service, By, WebDriverWait, EC, load_workbook, Workbook
    try:
        import customtkinter as ctk
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")
    except Exception as _ex:
        _IMPORTS_OK = False
        _IMPORTS_ERRO = f"customtkinter -- {_ex}"
        return

    _corrigir_dll_pywin32()
    try:
        import pandas as pd
        import win32com.client as win32
        from dotenv import load_dotenv
        from selenium import webdriver
        from selenium.webdriver.chrome.options import Options
        from selenium.webdriver.chrome.service import Service
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
        from openpyxl import load_workbook, Workbook
        load_dotenv()
        _IMPORTS_OK = True
        _IMPORTS_ERRO = ""
    except Exception as _ex:
        _IMPORTS_OK = False
        _IMPORTS_ERRO = str(_ex)

_tentar_importar_reais()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# CONFIGURAÇÕES
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# Chrome e driver portáteis de CAMINHO_REDE_BIN > o Chrome instalado se atualiza e fica incompatível com o driver.
_LOCAL_BASE     = os.path.join(os.environ.get("LOCALAPPDATA", pasta_script), "BureauMaestro")
PASTA_BIN       = os.path.join(_LOCAL_BASE, "Bin")
PASTA_PERFIL_ROBO = os.path.join(_LOCAL_BASE, "Perfil_Robo_Chrome")
CHROME_PATH     = os.path.join(PASTA_BIN, "chrome-win64",       "chrome.exe")
DRIVER_PATH     = os.path.join(PASTA_BIN, "chromedriver-win64", "chromedriver.exe")

LOGIN_EMAIL      = os.getenv("LOGIN_EMAIL",      "")
LOGIN_SENHA      = os.getenv("LOGIN_SENHA",      "")
CAMINHO_REDE     = os.getenv("CAMINHO_REDE",     "")
LINK_EMISSAO     = os.getenv("LINK_EMISSAO",     "")
LINK_ETIQUETAS   = os.getenv("LINK_ETIQUETAS",   "")
CAMINHO_REDE_BIN = os.getenv("CAMINHO_REDE_BIN", "")

CAMINHO_BACKUP = (
    os.path.join(os.path.dirname(CAMINHO_REDE), "BASE GERAL_BACKUP.xlsx")
    if CAMINHO_REDE else ""
)
PASTA_HISTORICO_EXCLUSOES = (
    os.path.join(os.path.dirname(CAMINHO_REDE), "HISTORICO DE EXCLUSOES")
    if CAMINHO_REDE else ""
)


NOME_ARQ_EMISSAO   = os.getenv("NOME_ARQ_EMISSAO",   "CONTROLE EMISSÃO 2026*")
NOME_ARQ_ETIQUETAS = os.getenv("NOME_ARQ_ETIQUETAS", "Solicitação de Etiquetas*")


CATEGORIA_PARCEIRO = os.getenv("CATEGORIA_PARCEIRO", "")

# MAPA_AGENDADOS: coluna de origem → coluna de destino; demais mapas: número da coluna de destino → coluna de origem.
MAPA_AGENDADOS = {
    "PEDIDO": "PEDIDO", "DATA AGENDA": "DATA", "ARTIGO": "ARTIGO",
    "DESCR ARTIGO": "DESCR", "AGENDADO DIA": "QTD",
    "VALIDACAO MODELO": "MODELO", "FORNECEDOR": "FORNECEDOR",
    "IMPRESSOS": "IMPRESSO", "BIPADO": "BIPADO"
}
MAPA_FORA = {
    2: "Pedido:", 3: "Artigo:", 4: "DATA_LIMPA", 6: "MOTIVO BUREAU",
    8: "Quantidade:", 9: "MODELO", 10: "IMPRESSOS"
}
MAPA_PARCEIRO = {
    2: "Pedido:", 3: "Artigo:", 4: "DATA_LIMPA",
    7: "Quantidade:", 8: "MODELO", 9: "IMPRESSOS"
}
MAPA_FORA_IMPRIMIR = {
    2: "Pedido:", 3: "Artigo:", 4: "DATA_LIMPA", 7: "Setor:",
    8: "Quantidade:", 9: "MODELO", 10: "IMPRESSOS"
}

if _IMPORTS_OK:
    EXCEL_EPOCH = pd.Timestamp("1899-12-30")

TEMPOS_RETRY  = [3 * 60, 5 * 60, 10 * 60]


# Pausa noturna: encerra às 20h e retoma às 8h, sem fechar o processo.
HORA_ENCERRAR = 20
HORA_RETOMAR  = 8


# Heartbeat a cada 10 min; lock expira após 2h sem atualização > PC travado não prende o lock.
INTERVALO_HEARTBEAT = 10 * 60
EXPIRACAO_LOCK       = 2 * 60 * 60

_status_atual = "Aguardando..."
_parar_loop   = threading.Event()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# PALETA VISUAL
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
COR_FUNDO        = "#18181B"
COR_CARD         = "#242428"
COR_HEADER       = "#FFC72C"
COR_HEADER_TXT   = "#1A1A1A"
COR_TEXTO        = "#F2F2F2"
COR_DETALHE      = "#A8A8A8"
COR_MUTED        = "#7A7A7A"
COR_BORDA        = "#3A3A3C"
COR_BTN_C        = "#333336"
COR_BTN_C_TXT    = "#F2F2F2"
COR_BTN_B        = "#3D8BFF"
COR_BTN_B_TXT    = "#FFFFFF"
COR_OK           = "#4ADE80"
COR_OK_FUNDO     = "#173321"
COR_AVISO        = "#FBBF24"
COR_AVISO_FUNDO  = "#3A2E0C"
COR_ERRO         = "#F87171"
COR_ERRO_FUNDO   = "#3A1414"


# Crédito > aparece no canto superior direito do cabeçalho.
NOME_CRIADOR        = os.getenv("NOME_CRIADOR",        "Rhebecca-Taih")
LINK_GITHUB_CRIADOR = os.getenv("LINK_GITHUB_CRIADOR", "https://github.com/Rhebecca-Taih")


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# ARQUIVOS DE CONTROLE
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_pasta_logs    = os.path.join(pasta_script, "logs")
_pasta_runtime = os.path.join(pasta_script, "runtime")
os.makedirs(_pasta_logs,    exist_ok=True)
os.makedirs(_pasta_runtime, exist_ok=True)

LOG_ARQUIVO          = os.path.join(_pasta_logs,    "historico_robo.txt")
LOG_MUDANCAS         = os.path.join(_pasta_logs,    "historico_mudancas.txt")
LOG_ERROS_PERMANENTE = os.path.join(_pasta_logs,    "historico_erros.txt")
LOCK_ARQUIVO         = os.path.join(_pasta_runtime, "bureau_maestro.lock")
MAQUINAS_ARQUIVO     = os.path.join(_pasta_runtime, "maquinas.txt")


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LOCK
# Formato do arquivo: PID|NOME_PC|DATA_INICIO|ULTIMO_HEARTBEAT
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _verificar_instancia_unica():
    nome_pc = os.environ.get("COMPUTERNAME", "DESCONHECIDO")
    if os.path.exists(LOCK_ARQUIVO):
        try:
            with open(LOCK_ARQUIVO, "r") as f:
                partes = f.read().strip().split("|")
            pid_antigo   = int(partes[0])
            pc_antigo    = partes[1] if len(partes) > 1 else "?"
            dt_inicio    = partes[2] if len(partes) > 2 else "?"
            dt_heartbeat = partes[3] if len(partes) > 3 else dt_inicio

            if pc_antigo == nome_pc:

                r = subprocess.run(f'tasklist /FI "PID eq {pid_antigo}" /NH',
                                   shell=True, capture_output=True, text=True,
                                   creationflags=_NO_WINDOW)
                if str(pid_antigo) in r.stdout:
                    return False, f"{pc_antigo} desde {dt_inicio}"

            else:


                try:
                    hb_dt = datetime.strptime(dt_heartbeat, "%d/%m/%Y %H:%M")
                    idade_seg = (datetime.now() - hb_dt).total_seconds()
                except Exception:
                    idade_seg = 0

                if idade_seg < EXPIRACAO_LOCK:
                    minutos_atras = max(0, int(idade_seg // 60))
                    return False, f"{pc_antigo} desde {dt_inicio} (última atividade há {minutos_atras} min)"

        except Exception:
            pass
        try:
            os.remove(LOCK_ARQUIVO)
        except Exception:
            pass
    try:
        agora_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        with open(LOCK_ARQUIVO, "w") as f:
            f.write(f"{os.getpid()}|{nome_pc}|{agora_str}|{agora_str}")
    except Exception:
        pass
    return True, ""

def _atualizar_heartbeat_lock():
    try:
        if not os.path.exists(LOCK_ARQUIVO):
            return
        with open(LOCK_ARQUIVO, "r") as f:
            partes = f.read().strip().split("|")
        if not partes or int(partes[0]) != os.getpid():
            return
        pc      = partes[1] if len(partes) > 1 else os.environ.get("COMPUTERNAME", "DESCONHECIDO")
        inicio  = partes[2] if len(partes) > 2 else datetime.now().strftime("%d/%m/%Y %H:%M")
        with open(LOCK_ARQUIVO, "w") as f:
            f.write(f"{os.getpid()}|{pc}|{inicio}|{datetime.now().strftime('%d/%m/%Y %H:%M')}")
    except Exception:
        pass

def _liberar_lock():
    try:
        if os.path.exists(LOCK_ARQUIVO):
            with open(LOCK_ARQUIVO, "r") as f:
                pid = int(f.read().strip().split("|")[0])
            if pid == os.getpid():
                os.remove(LOCK_ARQUIVO)
    except Exception:
        pass


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# MÁQUINAS
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _ler_maquinas():
    maquinas = {}
    if os.path.exists(MAQUINAS_ARQUIVO):
        try:
            with open(MAQUINAS_ARQUIVO, "r", encoding="utf-8") as f:
                for linha in f:
                    p = linha.strip().split("|")
                    if len(p) == 3:
                        maquinas[p[0]] = {"primeiro": p[1], "ultimo": p[2]}
        except Exception:
            pass
    return maquinas

def _registrar_maquina():
    nome_pc  = os.environ.get("COMPUTERNAME", "DESCONHECIDO")
    agora    = datetime.now().strftime("%d/%m/%Y %H:%M")
    maquinas = _ler_maquinas()
    eh_nova  = nome_pc not in maquinas
    if eh_nova:
        maquinas[nome_pc] = {"primeiro": agora, "ultimo": agora}
    else:
        maquinas[nome_pc]["ultimo"] = agora
    try:
        with open(MAQUINAS_ARQUIVO, "w", encoding="utf-8") as f:
            for nm, d in maquinas.items():
                f.write(f"{nm}|{d['primeiro']}|{d['ultimo']}\n")
    except Exception:
        pass
    return nome_pc, maquinas, eh_nova


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LOG
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _data_hoje():
    return datetime.now().strftime("%d/%m/%Y")

def _agora_completo():
    return datetime.now().strftime("%d/%m/%Y %H:%M:%S")

def _gerenciar_arquivo_log(caminho):
    sep = "=" * 60
    data_hoje = _data_hoje()
    if os.path.exists(caminho):
        try:
            with open(caminho, "r", encoding="utf-8") as f:
                f.readline()
                segunda = f.readline()
            if data_hoje not in segunda:
                os.remove(caminho)
        except Exception:
            try:
                os.remove(caminho)
            except Exception:
                pass
    if not os.path.exists(caminho):
        try:
            with open(caminho, "w", encoding="utf-8") as f:
                f.write(sep + "\n")
                f.write(f"  LOG DE AUTOMACAO -- {data_hoje}\n")
                f.write(sep + "\n\n")
        except Exception:
            pass

def _gerenciar_log():
    _gerenciar_arquivo_log(LOG_ARQUIVO)
    _gerenciar_arquivo_log(LOG_MUDANCAS)

PASTA_LOG_LOCAL = os.path.join(os.environ.get("LOCALAPPDATA") or os.path.expanduser("~"), "RobosBureau", "Base Geral")
DIAS_MANTER_LOG_LOCAL = 30
_rede_indisponivel_desde = None

def _limpar_logs_locais_antigos():
    try:
        limite = time.time() - DIAS_MANTER_LOG_LOCAL * 86400
        for nome in os.listdir(PASTA_LOG_LOCAL):
            caminho = os.path.join(PASTA_LOG_LOCAL, nome)
            if os.path.isfile(caminho) and os.path.getmtime(caminho) < limite:
                os.remove(caminho)
    except Exception:
        pass

def _registrar_log_local(texto):
    global _rede_indisponivel_desde
    if _rede_indisponivel_desde is None:
        _rede_indisponivel_desde = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    try:
        os.makedirs(PASTA_LOG_LOCAL, exist_ok=True)
        caminho = os.path.join(PASTA_LOG_LOCAL, f"log_{datetime.now().strftime('%d-%m-%Y')}.txt")
        arquivo_novo = not os.path.exists(caminho)
        with open(caminho, "a", encoding="utf-8") as f:
            f.write(texto)
        if arquivo_novo:
            _limpar_logs_locais_antigos()
    except Exception:
        pass

def _aviso_rede_restabelecida():
    if _rede_indisponivel_desde is None:
        return ""
    return (f"⚠️  Rede indisponível de {_rede_indisponivel_desde} até "
            f"{datetime.now().strftime('%d/%m/%Y %H:%M:%S')}. Registros desse período em: {PASTA_LOG_LOCAL}")

def _confirmar_rede_ok():
    global _rede_indisponivel_desde
    _rede_indisponivel_desde = None

def registrar_log(linhas):
    sep = "-" * 60
    texto_linhas = linhas if isinstance(linhas, list) else [str(linhas)]
    bloco = sep + "\n" + "".join(l + "\n" for l in texto_linhas) + "\n"
    try:
        with open(LOG_ARQUIVO, "a", encoding="utf-8") as f:
            aviso = _aviso_rede_restabelecida()
            if aviso:
                f.write(f"{sep}\n{aviso}\n\n")
            f.write(bloco)
        _confirmar_rede_ok()
    except Exception:
        _registrar_log_local(f"[{_agora_completo()}]\n{bloco}")

def registrar_mudanca(linhas):
    sep = "-" * 60
    texto_linhas = linhas if isinstance(linhas, list) else [str(linhas)]
    bloco = sep + "\n" + "".join(l + "\n" for l in texto_linhas) + "\n"
    try:
        with open(LOG_MUDANCAS, "a", encoding="utf-8") as f:
            f.write(bloco)
    except Exception:
        _registrar_log_local(f"[{_agora_completo()}] MUDANÇA\n{bloco}")

def registrar_erro_permanente(msg):
    try:
        with open(LOG_ERROS_PERMANENTE, "a", encoding="utf-8") as f:
            f.write(f"[{_agora_completo()}] {msg}\n")
    except Exception:
        _registrar_log_local(f"[{_agora_completo()}] ERRO: {msg}\n")

# Status novo precisa entrar na lista de palavras-chave > senão, não aparece como último status.
def _ultimo_status_log():
    if not os.path.exists(LOG_ARQUIVO):
        return "Sem registros anteriores."
    try:
        with open(LOG_ARQUIVO, "r", encoding="utf-8") as f:
            linhas = [l.strip() for l in f.readlines()
                      if l.strip() and not l.strip().startswith("-")
                      and not l.strip().startswith("=")
                      and "LOG DE AUTOMACAO" not in l]
        for linha in reversed(linhas):
            if any(e in linha for e in ["SUCESSO", "INICIADO", "AGUARDANDO", "ENCERRADO", "ERRO", "FALHA", "PAUSA", "RETOMANDO"]):
                return linha
        return linhas[-1] if linhas else "Sem registros."
    except Exception:
        return "Erro ao ler o log."


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# CHROME E BIN LOCAL
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _bin_completa():
    return os.path.exists(CHROME_PATH) and os.path.exists(DRIVER_PATH)

def _copiar_bin_da_rede():
    if _bin_completa():
        return True, ""
    if not CAMINHO_REDE_BIN:
        return False, "CAMINHO_REDE_BIN não configurado no .env"
    if not os.path.exists(CAMINHO_REDE_BIN):
        return False, f"Caminho de rede não encontrado: {CAMINHO_REDE_BIN}"
    try:
        os.makedirs(_LOCAL_BASE, exist_ok=True)
        shutil.copytree(CAMINHO_REDE_BIN, _LOCAL_BASE, dirs_exist_ok=True)
        if _bin_completa():
            return True, ""
        return False, "Cópia terminou mas os arquivos esperados não foram encontrados."
    except Exception as e:
        return False, str(e)

def _checar_chrome():
    r = {"chrome_existe": False, "driver_existe": False, "compat": False, "erro": ""}
    if not os.path.exists(CHROME_PATH):
        r["erro"] = (f"Chrome não encontrado.\n"
                     f"   Copie a pasta Bin para:\n"
                     f"   {PASTA_BIN}")
        return r
    r["chrome_existe"] = True
    if not os.path.exists(DRIVER_PATH):
        r["erro"] = (f"Chromedriver não encontrado.\n"
                     f"   Copie a pasta Bin para:\n"
                     f"   {PASTA_BIN}")
        return r
    r["driver_existe"] = True
    r["compat"] = True
    return r


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LINHAS DE STATUS ESTILIZADAS
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
TAG_ESTILO = {
    "secao":   {"tint": None,             "font": ("Segoe UI", 12, "bold"), "cor": COR_TEXTO,   "indent": 0,  "pady_top": 10},
    "ok":      {"tint": COR_OK_FUNDO,     "font": ("Segoe UI", 11),         "cor": COR_TEXTO,   "indent": 8,  "pady_top": 2},
    "aviso":   {"tint": COR_AVISO_FUNDO,  "font": ("Segoe UI", 11),         "cor": COR_TEXTO,   "indent": 8,  "pady_top": 2},
    "erro":    {"tint": COR_ERRO_FUNDO,   "font": ("Segoe UI", 11),         "cor": COR_TEXTO,   "indent": 8,  "pady_top": 2},
    "detalhe": {"tint": None,             "font": ("Segoe UI", 10),         "cor": COR_MUTED,   "indent": 20, "pady_top": 1},
    "normal":  {"tint": None,             "font": ("Segoe UI", 11),         "cor": COR_TEXTO,   "indent": 8,  "pady_top": 2},
}

def _criar_linha_log(container, texto, tag="normal", wraplength=460):
    if not texto or not texto.strip():
        ctk.CTkFrame(container, fg_color="transparent", height=6).pack(fill="x")
        return None, None
    estilo = TAG_ESTILO.get(tag, TAG_ESTILO["normal"])
    linha = ctk.CTkFrame(container, fg_color=(estilo["tint"] or "transparent"), corner_radius=8)
    linha.pack(fill="x", padx=4, pady=(estilo["pady_top"], 0))
    lbl = ctk.CTkLabel(linha, text=texto.strip(), font=estilo["font"], text_color=estilo["cor"],
                       anchor="w", justify="left", wraplength=wraplength)
    lbl.pack(anchor="w", padx=(estilo["indent"] + 8, 8), pady=6)
    return linha, lbl

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# JANELA DE LANÇAMENTO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
class JanelaLancamento:
    def __init__(self):
        self.aprovado      = False
        self._falha_lib    = False
        self._falha_chrome = False

        self.root = ctk.CTk()
        self.root.title("Base Geral")
        self.root.configure(fg_color=COR_FUNDO)
        self.root.resizable(False, False)
        self.root.attributes("-topmost", True)
        self.root.protocol("WM_DELETE_WINDOW", self._fechar)

        self._construir_ui()
        self._centralizar(560, 680)
        threading.Thread(target=self._verificar_lock_inicial, daemon=True).start()
        self.root.mainloop()

    def _centralizar(self, w, h):
        self.root.update_idletasks()
        sw = self.root.winfo_screenwidth()
        sh = self.root.winfo_screenheight()
        self.root.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")


    def _construir_ui(self):
        hdr = ctk.CTkFrame(self.root, fg_color="transparent")
        hdr.pack(fill="x", padx=24, pady=(24, 16))

        badge = ctk.CTkLabel(hdr, text="🤖", font=("Segoe UI", 20),
                             fg_color=COR_HEADER, text_color=COR_HEADER_TXT,
                             corner_radius=12, width=44, height=44)
        badge.pack(side="left")

        txt_hdr = ctk.CTkFrame(hdr, fg_color="transparent")
        txt_hdr.pack(side="left", padx=(12, 0))
        ctk.CTkLabel(txt_hdr, text="Base Geral", font=("Segoe UI", 17, "bold"),
                    text_color=COR_TEXTO, anchor="w").pack(anchor="w")
        ctk.CTkLabel(txt_hdr, text="Sincronizador de Base Geral", font=("Segoe UI", 12),
                    text_color=COR_DETALHE, anchor="w").pack(anchor="w")

        ctk.CTkButton(
            hdr, text=f"👤  CRIADO POR: {NOME_CRIADOR}", font=("Segoe UI", 9),
            fg_color="transparent", text_color=COR_MUTED, hover_color=COR_CARD,
            corner_radius=6, height=22, anchor="e",
            command=lambda: webbrowser.open(LINK_GITHUB_CRIADOR)
        ).pack(side="right", anchor="n")

        self.frm_controle = ctk.CTkFrame(self.root, fg_color="transparent")
        self.frm_controle.pack(fill="both", expand=True, padx=24)

        self.frm_btn = ctk.CTkFrame(self.root, fg_color="transparent")
        self.frm_btn.pack(fill="x", padx=24, pady=(12, 20))
        ctk.CTkButton(self.frm_btn, text="✖  Fechar", font=("Segoe UI", 11),
                     fg_color=COR_BTN_C, text_color=COR_BTN_C_TXT, hover_color=COR_BORDA,
                     corner_radius=8, height=36,
                     command=self._fechar).pack(side="right")

    def _limpar_controle(self):
        for w in self.frm_controle.winfo_children():
            w.destroy()

    def _card(self, titulo=None):
        card = ctk.CTkFrame(self.frm_controle, fg_color=COR_CARD, corner_radius=12,
                            border_width=1, border_color=COR_BORDA)
        card.pack(fill="x", pady=(0, 12))
        interno = ctk.CTkFrame(card, fg_color="transparent")
        interno.pack(fill="x", padx=16, pady=14)
        if titulo:
            ctk.CTkLabel(interno, text=titulo.upper(), font=("Segoe UI", 10, "bold"),
                        text_color=COR_MUTED, anchor="w").pack(anchor="w", pady=(0, 8))
        return interno

    def _ir_para_etapa(self, etapa):
        self._limpar_controle()
        getattr(self, f"_etapa_{etapa}")()

    def _mostrar_continuar(self, spoiler_texto, proxima_etapa, texto_botao="Continuar  →"):
        frm = ctk.CTkFrame(self.frm_controle, fg_color="transparent")
        frm.pack(fill="x", pady=(14, 0))
        if spoiler_texto:
            ctk.CTkLabel(frm, text=spoiler_texto, font=("Segoe UI", 10),
                        text_color=COR_MUTED, anchor="w").pack(anchor="w", pady=(0, 6))
        ctk.CTkButton(frm, text=texto_botao, font=("Segoe UI", 12, "bold"),
                     fg_color=COR_HEADER, text_color=COR_HEADER_TXT, hover_color="#E6B325",
                     corner_radius=8, height=38,
                     command=lambda: self._ir_para_etapa(proxima_etapa)).pack(fill="x")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 0 -- início
    # ─── ⋆⋅☆⋅⋆ ───
    def _verificar_lock_inicial(self):
        time.sleep(0.3)
        pode_rodar, info_lock = _verificar_instancia_unica()
        if not pode_rodar:
            self.root.after(0, lambda: self._tela_bloqueado(info_lock))
            return
        self.root.after(0, lambda: self._ir_para_etapa("bibliotecas"))

    def _tela_bloqueado(self, info_lock):
        self._limpar_controle()
        card = self._card()
        ctk.CTkLabel(card, text="⚠️  ROBÔ JÁ ATIVO EM OUTRA MÁQUINA", font=("Segoe UI", 12, "bold"),
                    text_color=COR_AVISO, anchor="w").pack(anchor="w")
        ctk.CTkLabel(card, text=f"Em execução em: {info_lock}", font=("Segoe UI", 10),
                    text_color=COR_MUTED, anchor="w", wraplength=460,
                    justify="left").pack(anchor="w", pady=(6, 0))
        ctk.CTkLabel(card, text="Feche esta janela.", font=("Segoe UI", 10),
                    text_color=COR_MUTED, anchor="w").pack(anchor="w")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 1 -- bibliotecas
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_bibliotecas(self):
        card = self._card("Bibliotecas")
        self._frm_bibliotecas = ctk.CTkFrame(card, fg_color="transparent")
        self._frm_bibliotecas.pack(fill="x")

        self._lbl_status_geral = ctk.CTkLabel(card, text="Verificando...", font=("Segoe UI", 10),
                                              text_color=COR_MUTED, anchor="w")
        self._lbl_status_geral.pack(fill="x", pady=(10, 4))

        self._barra_progresso = ctk.CTkProgressBar(card, height=6, corner_radius=4,
                                                    fg_color=COR_BORDA, progress_color=COR_HEADER)
        self._barra_progresso.set(0)
        self._barra_progresso.pack(fill="x")
        threading.Thread(target=self._checar_bibliotecas, daemon=True).start()

    def _status_geral(self, texto):
        try:
            self.root.after(0, lambda: self._lbl_status_geral.configure(text=texto))
        except Exception:
            pass

    def _fade_in_label(self, widget, cor_final, passos=10, intervalo=18):
        def _hex_rgb(c):
            c = c.lstrip("#")
            return tuple(int(c[i:i+2], 16) for i in (0, 2, 4))
        r1, g1, b1 = _hex_rgb(COR_CARD)
        r2, g2, b2 = _hex_rgb(cor_final)
        def _tick(i=0):
            if i > passos:
                try:
                    widget.configure(text_color=cor_final)
                except Exception:
                    pass
                return
            frac = i / passos
            cor = "#{:02X}{:02X}{:02X}".format(
                int(r1 + (r2 - r1) * frac), int(g1 + (g2 - g1) * frac), int(b1 + (b2 - b1) * frac))
            try:
                widget.configure(text_color=cor)
            except Exception:
                return
            self.root.after(intervalo, lambda: _tick(i + 1))
        _tick()

    def _linha_biblioteca(self, nome, estado="instalando"):
        resultado = {}
        evento = threading.Event()
        def _run():
            try:
                linha = ctk.CTkFrame(self._frm_bibliotecas, fg_color="transparent")
                linha.pack(fill="x")
                interno = ctk.CTkFrame(linha, fg_color="transparent")
                interno.pack(fill="x", pady=8)

                canvas_icone = tk.Canvas(interno, width=16, height=16, bg=COR_CARD,
                                         highlightthickness=0)
                canvas_icone.pack(side="left", padx=(0, 8))

                lbl_nome = ctk.CTkLabel(interno, text=nome, font=("Segoe UI", 11),
                                        text_color=COR_CARD)
                lbl_nome.pack(side="left")

                lbl_resultado = ctk.CTkLabel(interno, text="", font=("Segoe UI", 12, "bold"),
                                             text_color=COR_CARD)
                lbl_resultado.pack(side="right")

                ctk.CTkFrame(self._frm_bibliotecas, fg_color=COR_BORDA, height=1).pack(fill="x")

                resultado["canvas_icone"] = canvas_icone
                resultado["lbl_nome"] = lbl_nome
                resultado["lbl_resultado"] = lbl_resultado
                resultado["spinner_ativo"] = {"ativo": False}

                self._fade_in_label(lbl_nome, COR_TEXTO)
            except Exception:
                pass
            evento.set()
        try:
            self.root.after(0, _run)
        except Exception:
            evento.set()
        evento.wait(timeout=2)
        if resultado:
            self._definir_estado_linha(resultado, estado)
        return resultado

    def _girar_spinner(self, canvas_spinner, ativo_ref):
        angulo = [0]
        def _tick():
            if not ativo_ref["ativo"] or not canvas_spinner:
                return
            angulo[0] = (angulo[0] + 24) % 360
            try:
                canvas_spinner.itemconfig("arco", start=angulo[0])
            except Exception:
                return
            self.root.after(40, _tick)
        try:
            self.root.after(0, _tick)
        except Exception:
            pass

    def _definir_estado_linha(self, referencia, estado):
        def _run():
            try:
                canvas = referencia["canvas_icone"]
                referencia["spinner_ativo"]["ativo"] = False
                canvas.delete("all")

                if estado == "instalando":
                    canvas.create_arc(2, 2, 14, 14, start=0, extent=270, style="arc",
                                      outline=COR_HEADER, width=2, tags="arco")
                    referencia["spinner_ativo"]["ativo"] = True
                    self._girar_spinner(canvas, referencia["spinner_ativo"])
                else:
                    canvas.create_oval(2, 2, 14, 14, outline=COR_HEADER, width=2)

                if estado in ("ok", "instalado"):
                    referencia["lbl_resultado"].configure(text="✓", text_color=COR_OK)
                elif estado == "erro":
                    referencia["lbl_resultado"].configure(text="✕", text_color=COR_ERRO)
            except Exception:
                pass
        try:
            self.root.after(0, _run)
        except Exception:
            pass

    def _checar_bibliotecas(self):
        total = len(_STATUS_INSTALL)
        prontas = 0

        def _atualizar_progresso():
            try:
                self.root.after(0, lambda: self._barra_progresso.set(prontas / total if total else 1))
            except Exception:
                pass

        pendentes = []
        for nome_pip, status in _STATUS_INSTALL.items():
            if status == "ok":
                self._status_geral(f"Verificando {nome_pip}...")
                self._linha_biblioteca(nome_pip, "ok")
                prontas += 1
                _atualizar_progresso()
                time.sleep(0.35)
            else:
                pendentes.append(nome_pip)

        for nome_pip in pendentes:
            modulo = next((ni for ni, np in _LIBS if np == nome_pip), nome_pip)
            self._status_geral(f"Instalando {nome_pip}...")
            referencia = self._linha_biblioteca(nome_pip, "instalando")
            try:
                subprocess.run([sys.executable, "-m", "pip", "install", nome_pip, "--quiet"],
                               capture_output=True, timeout=180, creationflags=_NO_WINDOW)
                importlib.invalidate_caches()
                if nome_pip == "pywin32":
                    _corrigir_dll_pywin32()
                __import__(modulo)
                _STATUS_INSTALL[nome_pip] = "ok"
                self._definir_estado_linha(referencia, "instalado")
                self._status_geral(f"{nome_pip} instalado!")
            except Exception as e_lib:
                _STATUS_INSTALL[nome_pip] = "erro"
                self._definir_estado_linha(referencia, "erro")
                self._status_geral(f"Falha ao instalar {nome_pip}")
                registrar_erro_permanente(f"INSTALAR {nome_pip} -- {e_lib}")
            prontas += 1
            _atualizar_progresso()
            time.sleep(0.2)

        if not _IMPORTS_OK:
            _tentar_importar_reais()

        self._falha_lib = not _IMPORTS_OK
        self._status_geral("Tudo pronto!" if _IMPORTS_OK else "Há um problema a resolver.")
        self.root.after(500, self._mostrar_proximo_bibliotecas)

    def _mostrar_proximo_bibliotecas(self):
        if _IMPORTS_OK:
            self._mostrar_continuar("Próximo: Pasta Bin (Chrome/Driver)", "bin")
            return
        frm = ctk.CTkFrame(self.frm_controle, fg_color="transparent")
        frm.pack(fill="x", pady=(14, 0))
        ctk.CTkLabel(frm, text=f"❌  Erro: {_IMPORTS_ERRO[:150]}", font=("Segoe UI", 10),
                    text_color=COR_ERRO, wraplength=460, justify="left").pack(anchor="w", pady=(0, 8))
        ctk.CTkLabel(frm, text="Algumas bibliotecas (como pywin32) só funcionam de "
                              "verdade depois de reiniciar o programa.",
                    font=("Segoe UI", 9), text_color=COR_MUTED,
                    wraplength=460, justify="left").pack(anchor="w", pady=(0, 8))
        ctk.CTkButton(frm, text="🔄  Reiniciar", font=("Segoe UI", 11, "bold"),
                     fg_color=COR_BTN_C, text_color=COR_BTN_C_TXT, hover_color=COR_BORDA,
                     corner_radius=8, height=36,
                     command=self._reiniciar_programa).pack(fill="x")

    def _reiniciar_programa(self):
        try:
            subprocess.Popen([sys.executable, os.path.abspath(__file__)])
        except Exception as e:
            registrar_erro_permanente(f"REINICIAR -- {e}")
        finally:
            _liberar_lock()
            try:
                self.root.destroy()
            except Exception:
                pass
            os._exit(0)


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 2 -- bin (Chrome/Driver)
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_bin(self):
        card = self._card("Pasta Bin (Chrome/Driver locais)")
        self._lbl_bin_status = ctk.CTkLabel(card, text="Verificando...", font=("Segoe UI", 11),
                                            text_color=COR_TEXTO, anchor="w")
        self._lbl_bin_status.pack(anchor="w")
        self._lbl_bin_detalhe = ctk.CTkLabel(card, text="", font=("Segoe UI", 9),
                                             text_color=COR_MUTED, anchor="w",
                                             wraplength=460, justify="left")
        self._lbl_bin_detalhe.pack(anchor="w", pady=(4, 0))
        threading.Thread(target=self._checar_bin, daemon=True).start()

    def _checar_bin(self):
        if _bin_completa():
            self.root.after(0, lambda: self._lbl_bin_status.configure(
                text="✅  Já existe localmente.", text_color=COR_OK))
            self.root.after(300, lambda: self._mostrar_continuar("Próximo: Chrome e Driver", "chrome_driver"))
            return

        self.root.after(0, lambda: self._lbl_bin_status.configure(
            text="⏳  Copiando da rede...", text_color=COR_AVISO))
        self.root.after(0, lambda: self._lbl_bin_detalhe.configure(
            text=f"Origem: {CAMINHO_REDE_BIN}\nIsso pode demorar vários minutos, aguarde..."))

        parar_ticker = threading.Event()
        inicio_copia = time.time()

        def _ticker():
            while not parar_ticker.is_set():
                decorrido = int(time.time() - inicio_copia)
                try:
                    self.root.after(0, lambda d=decorrido: self._lbl_bin_status.configure(
                        text=f"⏳  Copiando Bin... {d//60}min {d%60}s"))
                except Exception:
                    return
                time.sleep(1)

        threading.Thread(target=_ticker, daemon=True).start()
        ok_copia, erro_copia = _copiar_bin_da_rede()
        parar_ticker.set()

        if ok_copia:
            self.root.after(0, lambda: self._lbl_bin_status.configure(
                text="✅  Bin copiada com sucesso.", text_color=COR_OK))
            self.root.after(0, lambda: self._lbl_bin_detalhe.configure(text=""))
        else:
            self.root.after(0, lambda: self._lbl_bin_status.configure(
                text="❌  Falha ao copiar Bin da rede.", text_color=COR_ERRO))
            self.root.after(0, lambda: self._lbl_bin_detalhe.configure(text=erro_copia[:200]))
            registrar_erro_permanente(f"CÓPIA BIN -- {erro_copia}")

        self.root.after(300, lambda: self._mostrar_continuar("Próximo: Chrome e Driver", "chrome_driver"))


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 3 -- Chrome e Driver
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_chrome_driver(self):
        card = self._card("Chrome e Driver")
        cr = _checar_chrome()

        linhas = []
        linhas.append(("✅  Chrome encontrado (local)", COR_OK) if cr["chrome_existe"]
                      else ("❌  Chrome não encontrado", COR_ERRO))
        linhas.append(("✅  Driver encontrado (local)", COR_OK) if cr["driver_existe"]
                      else ("❌  Driver não encontrado", COR_ERRO))
        if cr["compat"]:
            linhas.append(("✅  Compatíveis", COR_OK))

        for texto, cor in linhas:
            ctk.CTkLabel(card, text=texto, font=("Segoe UI", 11), text_color=cor,
                        anchor="w").pack(anchor="w", pady=2)

        if not cr["chrome_existe"] or not cr["driver_existe"]:
            for l in cr["erro"].split("\n"):
                if l.strip():
                    ctk.CTkLabel(card, text=l.strip(), font=("Segoe UI", 9),
                                text_color=COR_MUTED, anchor="w").pack(anchor="w")

        self._falha_chrome = not cr["compat"]
        self._mostrar_continuar("Próximo: Último status conhecido", "status")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 4 -- último status
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_status(self):
        card = self._card("Último status")
        ultimo = _ultimo_status_log()
        cor = (COR_OK if "SUCESSO" in ultimo else
               COR_ERRO if any(e in ultimo for e in ["ERRO", "FALHA", "CRITICO"]) else COR_TEXTO)
        ctk.CTkLabel(card, text=ultimo, font=("Segoe UI", 11), text_color=cor,
                    anchor="w", wraplength=460, justify="left").pack(anchor="w")
        self._mostrar_continuar("Próximo: Máquinas que já rodaram este robô", "maquinas")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 5 -- máquinas
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_maquinas(self):
        card = self._card("Máquinas")
        nome_pc, maquinas, eh_nova = _registrar_maquina()
        if eh_nova:
            ctk.CTkLabel(card, text=f"🆕  {nome_pc}  (primeira vez)", font=("Segoe UI", 11),
                        text_color=COR_AVISO, anchor="w").pack(anchor="w")
        else:
            ult = maquinas[nome_pc]["ultimo"]
            ctk.CTkLabel(card, text=f"🖥️  {nome_pc}  (último uso: {ult})", font=("Segoe UI", 11),
                        text_color=COR_TEXTO, anchor="w").pack(anchor="w")

        outras = [(nm, d) for nm, d in maquinas.items() if nm != nome_pc]
        if outras:
            ctk.CTkLabel(card, text="Outras máquinas:", font=("Segoe UI", 9),
                        text_color=COR_MUTED, anchor="w").pack(anchor="w", pady=(8, 2))
            for nm, d in outras[:5]:
                ctk.CTkLabel(card, text=f"•  {nm}  --  {d['ultimo']}", font=("Segoe UI", 9),
                            text_color=COR_MUTED, anchor="w").pack(anchor="w")

        self._mostrar_continuar("Próximo: Verificar a Base Geral", "integridade")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA 6 -- integridade da Base Geral
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_integridade(self):
        card = self._card("Integridade da Base Geral")

        if not CAMINHO_REDE or not _IMPORTS_OK:
            ctk.CTkLabel(card, text="—", font=("Segoe UI", 11),
                        text_color=COR_MUTED, anchor="w").pack(anchor="w")
            self._mostrar_continuar("Pronto para iniciar!", "final")
            return

        if not os.path.exists(CAMINHO_REDE):
            ctk.CTkLabel(card, text="⚠️  Arquivo não encontrado na rede.", font=("Segoe UI", 11),
                        text_color=COR_AVISO, anchor="w").pack(anchor="w")
            ctk.CTkLabel(card, text=CAMINHO_REDE, font=("Segoe UI", 9), text_color=COR_MUTED,
                        anchor="w", wraplength=460, justify="left").pack(anchor="w")
            self._mostrar_continuar("Pronto para iniciar!", "final")
            return

        try:
            wb_c = load_workbook(CAMINHO_REDE, read_only=True)
            wb_c.close()
            ctk.CTkLabel(card, text="✅  BASE GERAL.xlsx -- íntegra", font=("Segoe UI", 11),
                        text_color=COR_OK, anchor="w").pack(anchor="w")
            self._mostrar_continuar("Pronto para iniciar!", "final")
        except Exception as e_i:
            ctk.CTkLabel(card, text="❌  BASE GERAL.xlsx parece corrompida!", font=("Segoe UI", 11),
                        text_color=COR_ERRO, anchor="w").pack(anchor="w")
            ctk.CTkLabel(card, text=str(e_i)[:100], font=("Segoe UI", 9), text_color=COR_MUTED,
                        anchor="w", wraplength=460, justify="left").pack(anchor="w")
            bk_ok = bool(CAMINHO_BACKUP and os.path.exists(CAMINHO_BACKUP))
            ctk.CTkLabel(card, text=("💾  Backup disponível." if bk_ok else "⚠️  Sem backup."),
                        font=("Segoe UI", 10), text_color=COR_AVISO, anchor="w").pack(anchor="w", pady=(4, 0))
            registrar_erro_permanente(f"INTEGRIDADE -- {e_i}")
            self._oferecer_restauracao(bk_ok, self._falha_lib or self._falha_chrome)

    def _oferecer_restauracao(self, backup_disponivel, tem_erros):
        frm = ctk.CTkFrame(self.frm_controle, fg_color="transparent")
        frm.pack(fill="x", pady=(14, 0))
        if backup_disponivel:
            ctk.CTkButton(frm, text="🔄  Restaurar Backup", font=("Segoe UI", 12, "bold"),
                         fg_color=COR_BTN_B, text_color=COR_BTN_B_TXT, hover_color="#0D47A1",
                         corner_radius=8, height=38,
                         command=self._restaurar_e_continuar).pack(fill="x", pady=(0, 8))
        if not tem_erros:
            ctk.CTkButton(frm, text="▶  Continuar assim mesmo", font=("Segoe UI", 11),
                         fg_color=COR_AVISO, text_color="#FFFFFF", hover_color="#96690A",
                         corner_radius=8, height=36,
                         command=lambda: self._ir_para_etapa("final")).pack(fill="x")

    def _restaurar_e_continuar(self):
        self._limpar_controle()
        card = self._card()
        lbl = ctk.CTkLabel(card, text="⏳  Restaurando backup...", font=("Segoe UI", 11),
                           text_color=COR_AVISO, anchor="w")
        lbl.pack(anchor="w")
        try:
            shutil.copy2(CAMINHO_BACKUP, CAMINHO_REDE)
            lbl.configure(text="✅  Backup restaurado! O robô iniciará em instantes...", text_color=COR_OK)
            self.aprovado = True
            self.root.after(1500, self.root.destroy)
        except Exception as e:
            lbl.configure(text=f"❌  Falha: {e}", text_color=COR_ERRO)
            registrar_erro_permanente(f"RESTAURAR BACKUP -- {e}")


    # ─── ⋆⋅☆⋅⋆ ───
    # ETAPA FINAL -- iniciar
    # ─── ⋆⋅☆⋅⋆ ───
    def _etapa_final(self):
        card = self._card()
        falha = self._falha_lib or self._falha_chrome
        if falha:
            ctk.CTkLabel(card, text="⚠️  Alguns itens tiveram problema -- revise antes de iniciar.",
                        font=("Segoe UI", 11), text_color=COR_AVISO, anchor="w",
                        wraplength=460, justify="left").pack(anchor="w")
        else:
            ctk.CTkLabel(card, text="✅  Tudo certo! Pronto para iniciar.",
                        font=("Segoe UI", 13, "bold"), text_color=COR_OK, anchor="w").pack(anchor="w")

        frm_btns = ctk.CTkFrame(self.frm_controle, fg_color="transparent")
        frm_btns.pack(fill="x", pady=(6, 0))
        if not falha:
            ctk.CTkButton(frm_btns, text="▶  Iniciar", font=("Segoe UI", 12, "bold"),
                         fg_color=COR_HEADER, text_color=COR_HEADER_TXT, hover_color="#E6B325",
                         corner_radius=8, height=38,
                         command=self._iniciar).pack(fill="x", pady=(0, 8))
        ctk.CTkButton(frm_btns, text="🔑  Login Manual", font=("Segoe UI", 11),
                     fg_color=COR_BTN_B, text_color=COR_BTN_B_TXT, hover_color="#0D47A1",
                     corner_radius=8, height=36,
                     command=self._login_manual).pack(fill="x")

    def _login_manual(self):
        self._limpar_controle()
        card = self._card()
        lbl = ctk.CTkLabel(card, text="🔑  Abrindo Chrome para login manual...", font=("Segoe UI", 11),
                           text_color=COR_AVISO, anchor="w")
        lbl.pack(anchor="w")
        ctk.CTkLabel(card, text="1. Faça login e complete o autenticador\n"
                              "2. FECHE o Chrome quando terminar\n"
                              "3. Clique em Iniciar normalmente",
                    font=("Segoe UI", 9), text_color=COR_MUTED, anchor="w",
                    justify="left").pack(anchor="w", pady=(6, 0))

        def _abrir():
            try:
                url = LINK_EMISSAO.split("?")[0] if LINK_EMISSAO else "https://login.microsoftonline.com"
                proc = subprocess.Popen([
                    CHROME_PATH,
                    f"--user-data-dir={PASTA_PERFIL_ROBO}",
                    "--no-sandbox",
                    "--disable-dev-shm-usage",
                    "--no-first-run",
                    "--disable-sync",
                    "--disable-features=SyncDisabledWithNoNetwork",
                    url
                ])
                self.root.after(0, lambda: lbl.configure(
                    text="Chrome aberto. Faça login e feche.", text_color=COR_OK))
                proc.wait()
                self.root.after(0, self._etapa_final)
            except Exception as e:

                erro_txt = str(e)[:120]
                self.root.after(0, lambda: lbl.configure(
                    text=f"❌  Erro: {erro_txt}", text_color=COR_ERRO))

        threading.Thread(target=_abrir, daemon=True).start()

    def _iniciar(self):
        self.aprovado = True
        self.root.destroy()

    def _fechar(self):
        self.aprovado = False
        _liberar_lock()
        self.root.destroy()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LIMPEZA DE PROCESSOS
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# Busca por WMI > o driver não guarda o PID do Chrome de forma confiável.
def limpar_motor_robo():
    try:
        perfil_escapado = PASTA_PERFIL_ROBO.replace("'", "''")
        ps_cmd = (
            "$drivers = Get-CimInstance Win32_Process -Filter \"Name='chromedriver.exe'\"; "
            "foreach ($d in $drivers) { "
            "  $filhos = Get-CimInstance Win32_Process -Filter \"ParentProcessId=$($d.ProcessId) AND Name='chrome.exe'\"; "
            "  foreach ($c in $filhos) { "
            f"    if ($c.CommandLine -like '*{perfil_escapado}*') {{ Write-Output $d.ProcessId }} "
            "  } "
            "}"
        )
        resultado = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps_cmd],
            capture_output=True, text=True, timeout=20, creationflags=_NO_WINDOW
        )
        pids = {p.strip() for p in resultado.stdout.splitlines() if p.strip().isdigit()}
        for pid in pids:
            subprocess.run(f"taskkill /PID {pid} /F /T", shell=True,
                           capture_output=True, creationflags=_NO_WINDOW)
    except Exception:
        pass

# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# NORMALIZAÇÃO DE DADOS
# limpar_chave > maiúsculo, para comparar chaves; limpar_str > preserva a caixa.
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def limpar_chave(valor):
    s = str(valor).replace("\xa0", " ").strip().upper().replace("'","")
    if s.endswith(".0"): s = s[:-2]
    s = " ".join(s.split())
    return s if s not in ["NAN","NONE","NAT",""] else ""

def limpar_str(s):
    try:
        if pd.isna(s): return ""
    except Exception:
        pass
    val = str(s).replace("\xa0"," ").strip().replace("'","")
    if val.endswith(".0"): val = val[:-2]
    val = " ".join(val.split())
    return val if val.lower() not in ["nan","none","nat",""] else ""

def limpar_numero(val):
    try:
        if pd.isna(val): return ""
    except Exception:
        pass
    s = str(val).strip()
    if s.lower() in ["nan","none","nat",""]: return ""
    if s.endswith(".0"): s = s[:-2]
    try:
        return str(int(s))
    except Exception:
        return s

def _eh_quantidade(nome):
    n = str(nome).upper()
    return "QUANTIDADE" in n or "QTD" in n

def normalizar_valor(val, nome_coluna):
    if "DATA" in str(nome_coluna).upper(): return para_serial(val)
    if _eh_quantidade(nome_coluna): return limpar_numero(val)
    return limpar_str(val)

def para_serial(val):
    if val is None or val == "": return ""
    try:
        if isinstance(val, (pd.Timestamp, datetime)):
            ts = pd.Timestamp(val)
            if pd.isna(ts): return ""
            return (ts.normalize() - EXCEL_EPOCH).days
        # 20000–60000 > faixa de serial de data válida (~1954 a 2064).
        if isinstance(val, (int, float)):
            if 20000 < val < 60000: return int(val)
            return ""
        s = str(val).strip().replace("'","")
        if s.lower() in ["nan","nat","none",""]: return ""
        if s.replace(".","",1).isdigit():
            n = float(s)
            if 20000 < n < 60000: return int(n)
        ts = pd.to_datetime(s, dayfirst=True, errors="coerce")
        if pd.notna(ts): return (ts.normalize() - EXCEL_EPOCH).days
    except Exception:
        pass
    return ""

# Só a coluna de data vira número > as demais ficam texto para manter formatação e zero à esquerda.
def gravar_celula(ws, linha, col, valor, nome_coluna):
    if "DATA" in str(nome_coluna).upper() and isinstance(valor, int):
        ws.Cells(linha, col).Value = valor
    else:
        ws.Cells(linha, col).Value = str(valor) if valor != "" else ""


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LEITURA DE PLANILHAS
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def ler_excel_com_datas(caminho, sheet_name=None):
    kwargs = {"sheet_name": sheet_name} if sheet_name else {}
    df = pd.read_excel(caminho, **kwargs)
    df.columns = df.columns.str.strip()
    return df

def normalizar_df(df):
    for col in df.columns:
        if "DATA" in col.upper(): df[col] = df[col].apply(para_serial)
        elif _eh_quantidade(col): df[col] = df[col].apply(limpar_numero)
        else: df[col] = df[col].apply(limpar_str)
    return df

def preparar_etiquetas(arq_etq):
    wb_raw = load_workbook(arq_etq, data_only=True)
    ws_raw = wb_raw.active
    cabecalho = [
        str(cell.value).strip() if cell.value is not None else f"COL_{i}"
        for i, cell in enumerate(ws_raw[1])
    ]
    # Coluna de carimbo localizada pelo nome > a posição varia entre exportações do formulário.
    col_carimbo_idx = next(
        (i for i, c in enumerate(cabecalho)
         if "CARIMBO" in c.upper() or "DATA/HORA" in c.upper()), None
    )
    registros = []
    for row in ws_raw.iter_rows(min_row=2, values_only=True):
        registro = dict(zip(cabecalho, row))
        if col_carimbo_idx is not None:
            val = row[col_carimbo_idx]
            try:
                serial = (pd.Timestamp(val).normalize() - EXCEL_EPOCH).days
            except Exception:
                serial = ""
            registro["DATA_LIMPA"] = serial
        registros.append(registro)
    wb_raw.close()
    df = pd.DataFrame(registros)
    if not df.empty:
        for col in df.columns:
            if col == "DATA_LIMPA": continue
            elif _eh_quantidade(col): df[col] = df[col].apply(limpar_numero)
            else: df[col] = df[col].apply(limpar_str)
    return df


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# DOWNLOAD
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def executar_download(link, nome_esperado, precisa_login=False):
    base_nome    = nome_esperado.replace(".xlsx","").replace("*","").strip()
    padrao_busca = os.path.join(pasta_downloads, f"{base_nome}*.xlsx")

    for tentativa, tempo_espera in enumerate(TEMPOS_RETRY, start=1):
        minutos = tempo_espera // 60
        registrar_log([
            f"Download: {nome_esperado}",
            f"Data: {_agora_completo()}",
            f"Tentativa {tentativa}/{len(TEMPOS_RETRY)} -- até {minutos} min...",
        ])

        options = Options()
        options.add_argument("--headless=new")
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--no-first-run")
        options.add_argument("--disable-sync")
        options.add_argument("--disable-features=SyncDisabledWithNoNetwork")
        options.add_argument(f"user-data-dir={PASTA_PERFIL_ROBO}")
        options.add_argument("--remote-debugging-port=9222")
        options.page_load_strategy = "normal"
        options.add_experimental_option("prefs", {
            "download.default_directory": pasta_downloads,
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "safebrowsing.enabled": False
        })
        options.binary_location = CHROME_PATH

        driver     = None
        chrome_pid = None
        try:
            svc    = Service(executable_path=DRIVER_PATH)
            driver = webdriver.Chrome(service=svc, options=options)
            driver.set_page_load_timeout(180)


            try:
                chrome_pid = svc.process.pid
            except Exception:
                chrome_pid = None

            driver.execute_cdp_cmd("Page.setDownloadBehavior", {
                "behavior": "allow", "downloadPath": pasta_downloads
            })
            driver.get(link)
            time.sleep(10)

            if precisa_login:
                # idSIButton9 se repete em todas as etapas do login Microsoft > comportamento do próprio site.
                try:
                    espera = WebDriverWait(driver, 30)
                    email  = espera.until(EC.presence_of_element_located((By.NAME, "loginfmt")))
                    email.send_keys(LOGIN_EMAIL)
                    driver.find_element(By.ID, "idSIButton9").click()
                    time.sleep(5)
                    pwd = espera.until(EC.presence_of_element_located((By.NAME, "passwd")))
                    pwd.send_keys(LOGIN_SENHA)
                    driver.find_element(By.ID, "idSIButton9").click()
                    time.sleep(5)
                    try:
                        driver.find_element(By.ID, "idSIButton9").click()
                    except Exception:
                        pass
                    registrar_log(["Login automático realizado."])
                except Exception:
                    registrar_log(["Login já estava ativo no perfil (sessão manual)."])

                # download=1 na URL > baixa direto, sem a tela do portal.
                sep          = "&" if "?" in link else "?"
                url_download = link + sep + "download=1"
                driver.get(url_download)


            inicio = time.time()
            while time.time() - inicio < tempo_espera:
                busca = glob.glob(padrao_busca)
                reais = [f for f in busca
                         if not f.endswith(".crdownload") and not f.endswith(".tmp")]
                if reais:
                    final = max(reais, key=os.path.getmtime)
                    registrar_log([
                        f"Download concluído: {os.path.basename(final)}",
                        f"Data: {_agora_completo()}",
                    ])
                    return final
                time.sleep(3)

        except Exception as e:
            registrar_log([f"ERRO Chrome (tentativa {tentativa}): {str(e)[:150]}"])
            registrar_erro_permanente(f"DOWNLOAD [{nome_esperado}] t{tentativa} -- {str(e)[:150]}")
        finally:
            if driver:
                try:
                    driver.quit()
                except Exception:
                    pass
            # Encerra também pelo PID > Chrome órfão trava o perfil no próximo ciclo.
            if chrome_pid:
                try:
                    subprocess.run(f"taskkill /PID {chrome_pid} /F /T",
                                   shell=True, capture_output=True,
                                   creationflags=_NO_WINDOW)
                except Exception:
                    pass

        registrar_log([f"Arquivo não encontrado após {minutos} min (tentativa {tentativa})"])
        if tentativa < len(TEMPOS_RETRY):
            limpar_motor_robo()
            time.sleep(30)

    registrar_log([f"FALHA DEFINITIVA: {nome_esperado}", f"Data: {_agora_completo()}"])
    registrar_erro_permanente(f"DOWNLOAD FALHOU [{nome_esperado}]")
    return None


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# DIFF E HISTÓRICO DE EXCLUSÕES
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# Contador na chave > duplicata legítima não vira uma linha só.
def calcular_diff_etiquetas(df_origem, df_destino, map_colunas, nome_aba=""):
    lista_novos, lista_atu = [], []
    dados_dest   = {}
    tracker_dest = {}
    total_antes  = len(df_destino)

    idx_ped = next((k for k,v in map_colunas.items() if "Pedido" in str(v)), 2)
    idx_art = next((k for k,v in map_colunas.items() if "Artigo" in str(v)), 3)
    idx_dat = next((k for k,v in map_colunas.items() if "DATA_LIMPA" in str(v)), 4)

    if not df_destino.empty:
        for idx, row in df_destino.iterrows():
            try:
                ped = limpar_chave(row.iloc[idx_ped-1])
                art = limpar_chave(row.iloc[idx_art-1])
                dat = para_serial(row.iloc[idx_dat-1])
            except Exception:
                continue
            if ped and art:
                base_key = (ped, art, dat)
                c = tracker_dest.get(base_key, 0)
                tracker_dest[base_key] = c + 1
                dados_dest[(base_key+(c,))] = {"row_idx": idx+2, "data": row}

    tracker_orig   = {}
    chaves_vistas  = set()
    amostras_novas = []
    amostras_dest  = list(dados_dest.keys())[:3]

    for _, row in df_origem.iterrows():
        ped = limpar_chave(row.get("Pedido:", ""))
        art = limpar_chave(row.get("Artigo:", ""))
        dat = para_serial(row.get("DATA_LIMPA", ""))
        if not ped or not art: continue
        base_key = (ped, art, dat)
        c = tracker_orig.get(base_key, 0)
        tracker_orig[base_key] = c + 1
        chave_final = (base_key+(c,))
        chaves_vistas.add(chave_final)
        pacote = {col_num: normalizar_valor(row.get(col_nome,""), col_nome)
                  for col_num, col_nome in map_colunas.items()}

        if chave_final not in dados_dest:
            lista_novos.append(pacote)
            if len(amostras_novas) < 3:
                amostras_novas.append(f"NOVO {repr(ped)} | {repr(art)} | DATA={dat}")
            registrar_mudanca([f"NOVO [{nome_aba}] -- {_agora_completo()}",
                               f"Pedido: {ped} | Artigo: {art} | Data: {dat}"])
        else:
            info  = dados_dest[chave_final]
            mudou = False
            mudancas = {}
            detalhes = []
            for col_num, val_novo in pacote.items():
                try:
                    val_ant_raw = info["data"].iloc[col_num-1]
                except Exception:
                    val_ant_raw = ""
                val_ant = normalizar_valor(val_ant_raw, map_colunas.get(col_num,""))
                if val_novo != val_ant:
                    mudou = True
                    mudancas[col_num] = val_novo
                    detalhes.append(f"   {map_colunas.get(col_num,f'COL_{col_num}')}: {repr(val_ant)} -> {repr(val_novo)}")
            if mudou:
                lista_atu.append({"linha": info["row_idx"], "mudancas": mudancas})
                registrar_mudanca([f"ATUALIZACAO [{nome_aba}] -- {_agora_completo()}",
                                   f"Pedido: {ped} | Artigo: {art} | Data: {dat}"] + detalhes)

    linhas_deletar = []
    chaves_excluir = set(dados_dest.keys()) - chaves_vistas
    if chaves_excluir:
        linhas_historico = []
        for ch in chaves_excluir:
            ped_x, art_x, dat_x, _c = ch
            info_x = dados_dest[ch]
            linha = {"PEDIDO": ped_x, "ARTIGO": art_x, "DATA": dat_x}
            for col_num, nome in map_colunas.items():
                try:
                    linha[nome] = info_x["data"].iloc[col_num-1]
                except Exception:
                    linha[nome] = ""
            linhas_historico.append(linha)
        caminho_lista, caminho_base = _salvar_historico_exclusoes(linhas_historico, nome_aba)
        if caminho_lista:
            linhas_deletar = [dados_dest[ch]["row_idx"] for ch in chaves_excluir]
            for ch in chaves_excluir:
                ped_x, art_x, dat_x, _c = ch
                registrar_mudanca([f"EXCLUSAO [{nome_aba}] -- {_agora_completo()}",
                                   f"Pedido: {ped_x} | Artigo: {art_x} | "
                                   f"Não encontrado mais na Solicitação de Etiquetas de origem."])
            registrar_log([f"{nome_aba} -- Excluídos: {len(linhas_deletar)} (cópia salva em {caminho_lista}"
                           + (f", base completa em {caminho_base})" if caminho_base else ")")])
        else:
            registrar_log([f"⚠️  {len(chaves_excluir)} pedido(s) sumiram de [{nome_aba}] na origem, mas "
                           f"Falha ao salvar a cópia de segurança -- nenhuma exclusão será aplicada "
                           f"neste ciclo, por segurança."])
            registrar_erro_permanente(
                f"EXCLUSAO {nome_aba} -- falha ao salvar histórico, {len(chaves_excluir)} pedido(s) NÃO excluídos por segurança."
            )

    if amostras_novas and nome_aba:
        diag = [f"Diagnostico {nome_aba}:"] + amostras_novas
        if amostras_dest:
            diag.append("DESTINO (primeiros):")
            for k in amostras_dest:
                diag.append(f"  {repr(k[0])} | {repr(k[1])} | DATA={k[2]}")
        registrar_log(diag)

    return lista_novos, lista_atu, total_antes, linhas_deletar


def _pasta_historico_exclusoes_do_dia():
    agora = datetime.now()
    return os.path.join(PASTA_HISTORICO_EXCLUSOES, agora.strftime("%Y"),
                        agora.strftime("%m"), agora.strftime("%d"))

# O_EXCL > duas máquinas nunca reservam o mesmo Vn.
def _reservar_proxima_versao_exclusoes(pasta_dia):

    os.makedirs(pasta_dia, exist_ok=True)
    n = 1
    while True:
        candidato = os.path.join(pasta_dia, f"V{n}_PEDIDOS_EXCLUIDOS_CICLO.xlsx")
        try:
            fd = os.open(candidato, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.close(fd)
            return candidato, n
        except FileExistsError:
            n += 1

def _salvar_historico_exclusoes(linhas_excluidas, nome_aba):

    if not linhas_excluidas or not PASTA_HISTORICO_EXCLUSOES:
        return None, None
    try:
        pasta_dia = _pasta_historico_exclusoes_do_dia()
        caminho_lista, n = _reservar_proxima_versao_exclusoes(pasta_dia)
        wb = Workbook()
        ws = wb.active
        ws.title = nome_aba[:31] if nome_aba else "EXCLUIDOS"
        colunas = list(linhas_excluidas[0].keys())
        ws.cell(row=1, column=1, value="EXCLUIDO EM")
        for c, nome in enumerate(colunas, start=2):
            ws.cell(row=1, column=c, value=nome)
        agora_txt = _agora_completo()
        for r, linha in enumerate(linhas_excluidas, start=2):
            ws.cell(row=r, column=1, value=agora_txt)
            for c, nome in enumerate(colunas, start=2):
                ws.cell(row=r, column=c, value=linha.get(nome, ""))
        wb.save(caminho_lista)
    except Exception as e:
        registrar_erro_permanente(f"HISTÓRICO EXCLUSÕES (lista) -- {type(e).__name__}: {e}")
        return None, None

    caminho_base = os.path.join(pasta_dia, f"V{n}_BASE_ANTES_DA_EXCLUSAO.xlsx")
    try:
        if CAMINHO_REDE and os.path.exists(CAMINHO_REDE):
            shutil.copy2(CAMINHO_REDE, caminho_base)
        else:
            caminho_base = None
    except Exception as e:
        registrar_erro_permanente(f"HISTÓRICO EXCLUSÕES (base completa) -- {type(e).__name__}: {e}")
        caminho_base = None

    return caminho_lista, caminho_base


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# SINCRONIZAÇÃO PRINCIPAL
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def sincronizar_bureau_maestro():
    _gerenciar_log()
    limpar_motor_robo()

    arq_emis = executar_download(LINK_EMISSAO,   NOME_ARQ_EMISSAO, True)
    arq_etq  = executar_download(LINK_ETIQUETAS, NOME_ARQ_ETIQUETAS, False)

    if not arq_emis or not arq_etq:
        registrar_log(["Sincronização abortada.", f"Data: {_agora_completo()}"])
        return

    registrar_log(["Lendo e cruzando dados...", f"Data: {_agora_completo()}"])
    lista_novos_agen, lista_atu_agen = [], []
    linhas_deletar_agen = []

    try:
        df_emis = ler_excel_com_datas(arq_emis, sheet_name="BASE_DADOS")
        df_emis = normalizar_df(df_emis)
        registrar_log([f"CONTROLE EMISSÃO -- {len(df_emis)} linhas"])
    except Exception as e:
        registrar_log([f"ERRO ao ler CONTROLE EMISSÃO: {e}"])
        registrar_erro_permanente(f"LER EMISSÃO -- {e}")
        return

    try:
        df_dest_agen = ler_excel_com_datas(CAMINHO_REDE, sheet_name="PEDIDOS AGENDADOS")
        df_dest_agen = normalizar_df(df_dest_agen)
    except Exception as e:
        registrar_log([f"ERRO ao ler PEDIDOS AGENDADOS: {e}"])
        registrar_erro_permanente(f"LER AGENDADOS -- {e}")
        return

    try:
        df_dest_fora = ler_excel_com_datas(CAMINHO_REDE, sheet_name="PEDIDOS FORA DE AGENDA")
    except Exception as e:
        registrar_log([f"ERRO ao ler FORA DE AGENDA: {e}"])
        registrar_erro_permanente(f"LER FORA -- {e}")
        return

    try:
        df_dest_parceiro = ler_excel_com_datas(CAMINHO_REDE, sheet_name=CATEGORIA_PARCEIRO)
    except Exception as e:
        registrar_log([f"ERRO ao ler {CATEGORIA_PARCEIRO}: {e}"])
        registrar_erro_permanente(f"LER PARCEIRO -- {e}")
        return

    try:
        df_dest_fora_imprimir = ler_excel_com_datas(CAMINHO_REDE, sheet_name="PEDIDOS FORA IMPRIMIR")
    except Exception as e:
        registrar_log([f"⚠️  Aba PEDIDOS FORA IMPRIMIR não encontrada/erro ao ler -- pulando essa aba nesse ciclo: {e}"])
        registrar_erro_permanente(f"LER FORA IMPRIMIR -- {e}")
        df_dest_fora_imprimir = None

    try:
        total_antes  = len(df_dest_agen)
        dados_agen   = {}
        tracker_agen = {}

        for idx, row in df_dest_agen.iterrows():
            ped = limpar_chave(row.get("PEDIDO",""))
            art = limpar_chave(row.get("ARTIGO",""))
            dat = para_serial(row.get("DATA",""))
            if ped and art:
                bk = (ped, art, dat)
                c  = tracker_agen.get(bk, 0)
                tracker_agen[bk] = c + 1
                dados_agen[(bk+(c,))] = {"row_idx": idx+2, "data": row}

        tracker_orig = {}
        chaves_vistas_emissao = set()
        for _, row in df_emis.iterrows():
            ped = limpar_chave(row.get("PEDIDO",""))
            art = limpar_chave(row.get("ARTIGO",""))
            dat = para_serial(row.get("DATA AGENDA",""))
            if not ped or not art: continue
            bk = (ped, art, dat)
            c  = tracker_orig.get(bk, 0)
            tracker_orig[bk] = c + 1
            chave  = (bk+(c,))
            chaves_vistas_emissao.add(chave)
            pacote = {MAPA_AGENDADOS[col]: normalizar_valor(row[col], col)
                      for col in MAPA_AGENDADOS if col in df_emis.columns}

            if chave not in dados_agen:
                lista_novos_agen.append(pacote)
                registrar_mudanca([f"NOVO [AGENDADOS] -- {_agora_completo()}",
                                   f"Pedido: {ped} | Artigo: {art} | Data: {dat}"])
            else:
                info  = dados_agen[chave]
                mudou = False
                mudancas = {}
                detalhes = []
                for col_dest, v_novo in pacote.items():
                    v_ant = normalizar_valor(info["data"].get(col_dest,""), col_dest)
                    if v_novo != v_ant:
                        mudou = True
                        mudancas[col_dest] = v_novo
                        detalhes.append(f"   {col_dest}: {repr(v_ant)} -> {repr(v_novo)}")
                if mudou:
                    lista_atu_agen.append({"linha": info["row_idx"], "mudancas": mudancas})
                    registrar_mudanca([f"ATUALIZACAO [AGENDADOS] -- {_agora_completo()}",
                                       f"Pedido: {ped} | Artigo: {art}"] + detalhes)

        chaves_excluir = set(dados_agen.keys()) - chaves_vistas_emissao
        if chaves_excluir:
            linhas_historico = [dados_agen[ch]["data"].to_dict() for ch in chaves_excluir]
            caminho_hist, caminho_base_hist = _salvar_historico_exclusoes(linhas_historico, "PEDIDOS AGENDADOS")
            if caminho_hist:
                linhas_deletar_agen = [dados_agen[ch]["row_idx"] for ch in chaves_excluir]
                for ch in chaves_excluir:
                    d = dados_agen[ch]["data"]
                    registrar_mudanca([f"EXCLUSAO [AGENDADOS] -- {_agora_completo()}",
                                       f"Pedido: {d.get('PEDIDO','')} | Artigo: {d.get('ARTIGO','')} | "
                                       f"Não encontrado mais no CONTROLE EMISSÃO de origem."])
                registrar_log([f"PEDIDOS AGENDADOS -- Excluídos: {len(linhas_deletar_agen)} "
                               f"(cópia salva em {caminho_hist}"
                               + (f", base completa em {caminho_base_hist})" if caminho_base_hist else ")")])
            else:
                registrar_log([f"⚠️  {len(chaves_excluir)} pedido(s) sumiram do CONTROLE EMISSÃO, mas "
                               f"Falha ao salvar a cópia de segurança -- nenhuma exclusão será aplicada "
                               f"neste ciclo, por segurança."])
                registrar_erro_permanente(
                    f"EXCLUSAO AGENDADOS -- falha ao salvar histórico, {len(chaves_excluir)} pedido(s) NÃO excluídos por segurança."
                )

        registrar_log([f"PEDIDOS AGENDADOS -- Novos: {len(lista_novos_agen)} | Atualizados: {len(lista_atu_agen)} | Total: {total_antes+len(lista_novos_agen)-len(linhas_deletar_agen)}"])

    except Exception as e:
        registrar_log([f"ERRO ao processar AGENDADOS: {e}"])
        registrar_erro_permanente(f"PROCESSAR AGENDADOS -- {e}")

    novos_fora, atu_fora, t_fora, deletar_fora = [], [], 0, []
    novos_parceiro, atu_parceiro, t_parceiro, deletar_parceiro = [], [], 0, []
    novos_fora_imp, atu_fora_imp, t_fora_imp, deletar_fora_imp = [], [], 0, []

    try:
        df_etq = preparar_etiquetas(arq_etq)
        df_fora = df_etq[df_etq.get("MOTIVO BUREAU", pd.Series(dtype=str))
                         .isin(["FORNECEDOR","NOVO","NACIONAL","IMPORTADOS","CDA","ERRO DE CADASTRO"])].copy()
        df_parceiro = df_etq[df_etq.get("MOTIVO BUREAU", pd.Series(dtype=str)) == CATEGORIA_PARCEIRO].copy()
        registrar_log([f"ETIQUETAS -- Total: {len(df_etq)} | FORA: {len(df_fora)} | {CATEGORIA_PARCEIRO}: {len(df_parceiro)}"])
        novos_fora, atu_fora, t_fora, deletar_fora = calcular_diff_etiquetas(df_fora, df_dest_fora, MAPA_FORA, "FORA DE AGENDA")
        novos_parceiro, atu_parceiro, t_parceiro, deletar_parceiro = calcular_diff_etiquetas(df_parceiro, df_dest_parceiro, MAPA_PARCEIRO, CATEGORIA_PARCEIRO)
        registrar_log([f"FORA DE AGENDA -- Novos: {len(novos_fora)} | Atualizados: {len(atu_fora)} | Excluídos: {len(deletar_fora)} | Total: {t_fora+len(novos_fora)-len(deletar_fora)}"])
        registrar_log([f"{CATEGORIA_PARCEIRO} -- Novos: {len(novos_parceiro)} | Atualizados: {len(atu_parceiro)} | Excluídos: {len(deletar_parceiro)} | Total: {t_parceiro+len(novos_parceiro)-len(deletar_parceiro)}"])

        if df_dest_fora_imprimir is not None:
            try:
                df_fora_imp = df_etq[df_etq.get("IMPRESSOS", pd.Series(dtype=str)) == ""].copy()
                registrar_log([f"FORA IMPRIMIR -- candidatos (IMPRESSOS vazio): {len(df_fora_imp)}"])
                novos_fora_imp, atu_fora_imp, t_fora_imp, deletar_fora_imp = calcular_diff_etiquetas(
                    df_fora_imp, df_dest_fora_imprimir, MAPA_FORA_IMPRIMIR, "FORA IMPRIMIR")
                registrar_log([f"FORA IMPRIMIR -- Novos: {len(novos_fora_imp)} | Atualizados: {len(atu_fora_imp)} | "
                               f"Excluídos: {len(deletar_fora_imp)} | Total: {t_fora_imp+len(novos_fora_imp)-len(deletar_fora_imp)}"])
            except Exception as e:
                registrar_log([f"ERRO ao processar FORA IMPRIMIR: {e}"])
                registrar_erro_permanente(f"PROCESSAR FORA IMPRIMIR -- {e}")
    except Exception as e:
        registrar_log([f"ERRO ao processar ETIQUETAS: {e}"])
        registrar_erro_permanente(f"PROCESSAR ETIQUETAS -- {e}")


    if CAMINHO_BACKUP:
        try:
            shutil.copy2(CAMINHO_REDE, CAMINHO_BACKUP)
            registrar_log(["Backup criado."])
        except Exception as e:
            registrar_erro_permanente(f"BACKUP -- {e}")

    # Edita cópia local e só substitui a de rede no fim > queda no meio não corrompe a base.
    temp_path = os.path.join(os.environ["TEMP"], "BUREAU_TEMP.xlsx")
    shutil.copy2(CAMINHO_REDE, temp_path)

    excel_pid = None
    excel = win32.DispatchEx("Excel.Application")
    try:
        excel_pid = excel.ProcessID
    except Exception:
        pass
    excel.Visible        = False
    excel.DisplayAlerts  = False
    excel.ScreenUpdating = False
    wb = excel.Workbooks.Open(temp_path)

    try:
        if lista_novos_agen or lista_atu_agen or linhas_deletar_agen:
            ws_agen = wb.Sheets("PEDIDOS AGENDADOS")
            # AutoFilter desligado antes de mexer nas linhas > filtro ativo desloca o índice das linhas.
            if ws_agen.AutoFilterMode: ws_agen.AutoFilterMode = False
            # Cabeçalho lido pelo nome na linha 1 > colunas reordenadas continuam funcionando.
            header_map = {
                str(ws_agen.Cells(1,c).Value).strip().upper(): c
                for c in range(1, ws_agen.UsedRange.Columns.Count+1)
                if ws_agen.Cells(1,c).Value
            }
            for item in lista_atu_agen:
                for col_nome, val in item["mudancas"].items():
                    if col_nome.upper() in header_map:
                        gravar_celula(ws_agen, int(item["linha"]), header_map[col_nome.upper()], val, col_nome)
            if lista_novos_agen:
                col_ped = header_map.get("PEDIDO", 2)
                # -4162 = xlUp > evita importar win32com.client.constants.
                prox = ws_agen.Cells(ws_agen.Rows.Count, col_ped).End(-4162).Row + 1
                if prox < 2: prox = 2
                for i, item in enumerate(lista_novos_agen):
                    for col_nome, val in item.items():
                        if col_nome.upper() in header_map:
                            gravar_celula(ws_agen, prox+i, header_map[col_nome.upper()], val, col_nome)
            if linhas_deletar_agen:
                # Exclui de baixo para cima > excluir de cima desloca o índice das linhas seguintes.
                for linha_num in sorted(set(linhas_deletar_agen), reverse=True):
                    ws_agen.Rows(int(linha_num)).Delete()

        def aplicar(aba, lista_atu, lista_novos, mapa, lista_deletar=None):
            ws = wb.Sheets(aba)
            # AutoFilter desligado antes de mexer nas linhas > filtro ativo desloca o índice das linhas.
            if ws.AutoFilterMode: ws.AutoFilterMode = False
            for item in lista_atu:
                for col_num, val in item["mudancas"].items():
                    gravar_celula(ws, int(item["linha"]), int(col_num), val, mapa.get(col_num,""))
            if lista_novos:
                col_ped = next((k for k,v in mapa.items() if "PEDIDO" in str(v).upper()), 2)
                # -4162 = xlUp > evita importar win32com.client.constants.
                prox = ws.Cells(ws.Rows.Count, col_ped).End(-4162).Row + 1
                if prox < 2: prox = 2
                for i, item in enumerate(lista_novos):
                    for col_num, val in item.items():
                        gravar_celula(ws, prox+i, int(col_num), val, mapa.get(col_num,""))
            if lista_deletar:
                # Exclui de baixo para cima > excluir de cima desloca o índice das linhas seguintes.
                lo_principal = next((lo for lo in ws.ListObjects if lo.Range.Column == 1), None)
                if lo_principal:
                    ultima_col = lo_principal.Range.Column + lo_principal.Range.Columns.Count - 1
                    for linha_num in sorted(set(lista_deletar), reverse=True):
                        ws.Range(ws.Cells(int(linha_num), 1),
                                 ws.Cells(int(linha_num), ultima_col)).Delete(Shift=-4162)
                else:
                    for linha_num in sorted(set(lista_deletar), reverse=True):
                        ws.Rows(int(linha_num)).Delete()

        aplicar("PEDIDOS FORA DE AGENDA", atu_fora, novos_fora, MAPA_FORA, deletar_fora)
        aplicar(CATEGORIA_PARCEIRO, atu_parceiro, novos_parceiro, MAPA_PARCEIRO, deletar_parceiro)

        if novos_fora_imp or atu_fora_imp or deletar_fora_imp:
            try:
                aplicar("PEDIDOS FORA IMPRIMIR", atu_fora_imp, novos_fora_imp, MAPA_FORA_IMPRIMIR, deletar_fora_imp)
            except Exception as e:
                registrar_log([f"⚠️  Falha ao gravar PEDIDOS FORA IMPRIMIR (aba existe na BASE GERAL?): {e}"])
                registrar_erro_permanente(f"GRAVAÇÃO FORA IMPRIMIR -- {e}")

        excel.ScreenUpdating = True
        wb.Save()
        registrar_log(["BASE GERAL SINCRONIZADA COM SUCESSO", f"Data: {_agora_completo()}"])

    except Exception as e:
        registrar_log([f"ERRO ao gravar Excel: {e}"])
        registrar_erro_permanente(f"GRAVACAO -- {e}")

    finally:
        try: wb.Close()
        except Exception: pass
        try: excel.Quit()
        except Exception: pass

        if excel_pid:
            try:
                subprocess.run(f"taskkill /PID {excel_pid} /F", shell=True,
                               capture_output=True, creationflags=_NO_WINDOW)
            except Exception:
                pass
        try:
            shutil.copy2(temp_path, CAMINHO_REDE)
        except Exception as e:
            registrar_erro_permanente(f"COPIA->REDE -- {e}")
        for f in [arq_emis, arq_etq, temp_path]:
            if f and os.path.exists(f):
                try: os.remove(f)
                except Exception: pass


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# LOOP PRINCIPAL
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
def _criar_icone_tray():
    try:
        from PIL import Image, ImageDraw
        img  = Image.new("RGBA", (64,64), (0,0,0,0))
        draw = ImageDraw.Draw(img)
        draw.ellipse([0,0,63,63],   fill="#18181B")
        draw.ellipse([16,16,47,47], fill="#FFC72C")
        draw.ellipse([24,24,39,39], fill="#18181B")
        return img
    except Exception:
        from PIL import Image
        return Image.new("RGBA", (64,64), "#18181B")

def _esta_em_pausa(agora=None):
    agora = agora or datetime.now()
    h = agora.hour
    if HORA_ENCERRAR <= HORA_RETOMAR:
        return HORA_ENCERRAR <= h < HORA_RETOMAR
    return h >= HORA_ENCERRAR or h < HORA_RETOMAR

def _loop_sincronizacao(forcar_primeiro_ciclo=False):
    while not _parar_loop.is_set():
        try:
            _loop_sincronizacao_interno(forcar_primeiro_ciclo=forcar_primeiro_ciclo)
            break
        except Exception as e:
            try:
                import traceback
                registrar_erro_permanente(f"LOOP MORREU SEM TRATAMENTO -- reiniciando automaticamente. {e}\n{traceback.format_exc()}")
                registrar_log([
                    "💥  O loop principal parou por um erro inesperado -- reiniciando automaticamente",
                    f"🕐  {_agora_completo()}",
                    f"     Detalhe: {e}",
                ])
            except Exception:
                pass
            forcar_primeiro_ciclo = False
            for _ in range(30):
                if _parar_loop.is_set(): break
                time.sleep(1)
    _liberar_lock()

def _loop_sincronizacao_interno(forcar_primeiro_ciclo=False):
    global _status_atual
    pythoncom = None
    try:
        import pythoncom
        pythoncom.CoInitialize()
    except Exception:
        pythoncom = None
    try:
        _gerenciar_log()
        registrar_log([
            "BASE GERAL -- SISTEMA INICIADO",
            f"Data: {_agora_completo()}",
            f"Base de rede: {CAMINHO_REDE}",
            f"PID: {os.getpid()} | PC: {os.environ.get('COMPUTERNAME','?')}",
            f"Pausa noturna: {HORA_ENCERRAR}h -- {HORA_RETOMAR}h (retomada automática)",
        ])
        _status_atual = "Ativo"

        ultimo_heartbeat = time.time()
        estava_pausado   = False
        primeiro_ciclo   = True

        while not _parar_loop.is_set():
            agora    = datetime.now()
            em_pausa = _esta_em_pausa(agora)


            ignorar_pausa_agora = em_pausa and primeiro_ciclo and forcar_primeiro_ciclo
            if ignorar_pausa_agora:
                registrar_log([
                    f"INÍCIO MANUAL dentro da janela de pausa ({HORA_ENCERRAR}h-{HORA_RETOMAR}h).",
                    f"Data: {_agora_completo()}",
                    "Ciclo imediato em execução antes de voltar à pausa.",
                ])


            if em_pausa and not ignorar_pausa_agora:
                if not estava_pausado:
                    registrar_log([
                        f"PAUSA -- Encerramento automático às {HORA_ENCERRAR}h.",
                        f"Data: {_agora_completo()}",
                        f"Retoma automaticamente às {HORA_RETOMAR}h.",
                    ])
                    estava_pausado = True
                _status_atual = f"🌙 Pausado até {HORA_RETOMAR:02d}:00"

                if time.time() - ultimo_heartbeat >= INTERVALO_HEARTBEAT:
                    _atualizar_heartbeat_lock()
                    ultimo_heartbeat = time.time()

                primeiro_ciclo = False
                for _ in range(60):
                    if _parar_loop.is_set(): break
                    time.sleep(1)
                continue

            primeiro_ciclo = False


            if estava_pausado:
                registrar_log([
                    f"RETOMANDO -- Operação normal retomada às {HORA_RETOMAR}h.",
                    f"Data: {_agora_completo()}",
                ])
                estava_pausado = False
                _status_atual = "Ativo"


            try:
                _status_atual = "Sincronizando..."
                sincronizar_bureau_maestro()
                _status_atual = f"OK -- {datetime.now().strftime('%H:%M')}"
                registrar_log(["Aguardando 15 min...", f"Data: {_agora_completo()}"])
            except Exception as e:
                _status_atual = f"Erro -- {datetime.now().strftime('%H:%M')}"
                registrar_log([f"ERRO CRÍTICO: {e}"])
                registrar_erro_permanente(f"LOOP -- {e}")

            _atualizar_heartbeat_lock()
            ultimo_heartbeat = time.time()

            for _ in range(180):
                if _parar_loop.is_set(): break
                if time.time() - ultimo_heartbeat >= INTERVALO_HEARTBEAT:
                    _atualizar_heartbeat_lock()
                    ultimo_heartbeat = time.time()
                time.sleep(5)

    finally:
        if pythoncom:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# HISTÓRICO AO VIVO
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
_historico_aberto = False

MAX_BLOCOS_HISTORICO   = 200
TAMANHO_LEITURA_INICIAL = 300_000

def _classificar_bloco(linhas):
    texto = " ".join(linhas).upper()
    if "SUCESSO" in texto:
        return "ok"
    if any(e in texto for e in ["ERRO", "FALHA", "CRITICO"]):
        return "erro"
    if any(e in texto for e in ["AVISO", "PAUSA", "ATENCAO"]):
        return "aviso"
    return "normal"

def _parsear_blocos_log(texto_bruto):
    sep = "-" * 60
    blocos = []
    for pedaco in texto_bruto.split(sep):
        linhas = [l.rstrip() for l in pedaco.strip("\n").split("\n") if l.strip()]
        if linhas:
            blocos.append(linhas)
    return blocos

class JanelaHistorico:
    def __init__(self):
        self.root = ctk.CTk()
        self.root.title("Base Geral -- Histórico")
        self.root.configure(fg_color=COR_FUNDO)
        self.root.resizable(False, False)
        self.root.geometry("580x680")
        self.root.protocol("WM_DELETE_WINDOW", self._fechar)

        self._parar   = threading.Event()
        self._pos_arquivo = 0

        self._construir_ui()
        self._carregar_inicial()
        threading.Thread(target=self._loop_tail, daemon=True).start()
        self.root.mainloop()

    def _construir_ui(self):
        hdr = ctk.CTkFrame(self.root, fg_color="transparent")
        hdr.pack(fill="x", padx=24, pady=(24, 12))
        badge = ctk.CTkLabel(hdr, text="📊", font=("Segoe UI", 20),
                             fg_color=COR_HEADER, text_color=COR_HEADER_TXT,
                             corner_radius=12, width=44, height=44)
        badge.pack(side="left")
        txt_hdr = ctk.CTkFrame(hdr, fg_color="transparent")
        txt_hdr.pack(side="left", padx=(12, 0))
        ctk.CTkLabel(txt_hdr, text="Histórico ao vivo", font=("Segoe UI", 16, "bold"),
                    text_color=COR_TEXTO, anchor="w").pack(anchor="w")
        ctk.CTkLabel(txt_hdr, text="Base Geral", font=("Segoe UI", 11),
                    text_color=COR_DETALHE, anchor="w").pack(anchor="w")
        ctk.CTkLabel(hdr, text="●  ao vivo", font=("Segoe UI", 10),
                    text_color=COR_OK).pack(side="right", anchor="ne", pady=4)

        self.frm_log = ctk.CTkScrollableFrame(self.root, fg_color=COR_CARD, corner_radius=12,
                                              border_width=1, border_color=COR_BORDA)
        self.frm_log.pack(fill="both", expand=True, padx=24, pady=(0, 12))

        rodape = ctk.CTkFrame(self.root, fg_color="transparent")
        rodape.pack(fill="x", padx=24, pady=(0, 20))
        ctk.CTkButton(rodape, text="📄  Abrir arquivo bruto (.txt)", font=("Segoe UI", 10),
                     fg_color=COR_BTN_C, text_color=COR_BTN_C_TXT, hover_color=COR_BORDA,
                     corner_radius=8, height=32,
                     command=lambda: os.startfile(LOG_ARQUIVO)).pack(side="left")
        ctk.CTkButton(rodape, text="✖  Fechar", font=("Segoe UI", 10),
                     fg_color=COR_BTN_C, text_color=COR_BTN_C_TXT, hover_color=COR_BORDA,
                     corner_radius=8, height=32,
                     command=self._fechar).pack(side="right")

    def _exibir_bloco(self, linhas):
        tag = _classificar_bloco(linhas)
        cabecalho, *resto = linhas
        _criar_linha_log(self.frm_log, cabecalho, tag)
        for l in resto:
            _criar_linha_log(self.frm_log, l, "detalhe")
        try:
            self.root.after(50, lambda: self.frm_log._parent_canvas.yview_moveto(1.0))
        except Exception:
            pass

    def _limpar_tela(self):
        for w in self.frm_log.winfo_children():
            w.destroy()

    def _carregar_inicial(self):
        try:
            with open(LOG_ARQUIVO, "rb") as f:
                f.seek(0, os.SEEK_END)
                tamanho = f.tell()
                ler = min(tamanho, TAMANHO_LEITURA_INICIAL)
                f.seek(tamanho - ler)
                bruto = f.read().decode("utf-8", errors="ignore")
                self._pos_arquivo = tamanho
        except Exception:
            bruto = ""
            self._pos_arquivo = 0

        blocos = _parsear_blocos_log(bruto)[-MAX_BLOCOS_HISTORICO:]
        if not blocos:
            _criar_linha_log(self.frm_log, "Sem registros ainda hoje.", "detalhe")
        for bloco in blocos:
            self._exibir_bloco(bloco)

    def _loop_tail(self):
        buffer_pendente = ""
        while not self._parar.is_set():
            try:
                tamanho_atual = os.path.getsize(LOG_ARQUIVO)
                if tamanho_atual < self._pos_arquivo:

                    self._pos_arquivo = 0
                    buffer_pendente = ""
                    self.root.after(0, self._limpar_tela)

                if tamanho_atual > self._pos_arquivo:
                    with open(LOG_ARQUIVO, "r", encoding="utf-8", errors="ignore") as f:
                        f.seek(self._pos_arquivo)
                        novo = f.read()
                    self._pos_arquivo = tamanho_atual
                    buffer_pendente += novo


                    sep = "-" * 60
                    partes = buffer_pendente.split(sep)
                    completos, buffer_pendente = partes[:-1], sep + partes[-1]
                    for pedaco in completos:
                        linhas = [l.rstrip() for l in pedaco.strip("\n").split("\n") if l.strip()]
                        if linhas:
                            self.root.after(0, lambda ls=linhas: self._exibir_bloco(ls))
            except Exception:
                pass
            time.sleep(2)

    def _fechar(self):
        global _historico_aberto
        self._parar.set()
        _historico_aberto = False
        try:
            self.root.destroy()
        except Exception:
            pass

def _abrir_historico_ao_vivo():
    global _historico_aberto
    if _historico_aberto:
        return
    _historico_aberto = True
    threading.Thread(target=_rodar_janela_historico, daemon=True).start()

def _rodar_janela_historico():
    global _historico_aberto
    try:
        JanelaHistorico()
    except Exception as e:
        registrar_erro_permanente(f"JANELA HISTÓRICO -- {e}")
    finally:
        _historico_aberto = False

def _iniciar_tray(forcar_primeiro_ciclo=False):
    try:
        import pystray
    except ImportError:
        _loop_sincronizacao(forcar_primeiro_ciclo=forcar_primeiro_ciclo)
        return

    def texto_status(item):
        return f"Status: {_status_atual}"

    def parar(icon, item):
        _parar_loop.set()


        _liberar_lock()
        registrar_log(["Robô parado pelo usuário.", f"Data: {_agora_completo()}"])
        threading.Timer(5, os._exit, args=(0,)).start()
        icon.stop()

    def abrir_status(icon, item):
        _abrir_historico_ao_vivo()

    menu = pystray.Menu(
        pystray.MenuItem("📊  Ver histórico", abrir_status, default=True),
        pystray.MenuItem(texto_status, None, enabled=False),
        pystray.Menu.SEPARATOR,
        pystray.MenuItem("⏹  Parar robô", parar),
    )
    icone = pystray.Icon("BaseGeral", _criar_icone_tray(), "Base Geral", menu=menu)

    threading.Thread(target=_loop_sincronizacao,
                     kwargs={"forcar_primeiro_ciclo": forcar_primeiro_ciclo},
                     daemon=True).start()
    icone.run()


# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
# PONTO DE ENTRADA
# ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆ ⋆⭒˚.⋆
if __name__ == "__main__":
    try:
        janela = JanelaLancamento()
    except Exception as _e:


        import traceback
        _erro_completo = traceback.format_exc()
        try:
            with open(os.path.join(pasta_script, "erro_fatal.txt"), "w", encoding="utf-8") as _f:
                _f.write(f"[{datetime.now().strftime('%d/%m/%Y %H:%M:%S')}]\n{_erro_completo}")
        except Exception:
            pass
        try:
            import tkinter.messagebox as _mb
            _raiz_erro = tk.Tk()
            _raiz_erro.withdraw()
            _mb.showerror(
                "Base Geral -- Erro ao iniciar",
                f"O robô encontrou um erro e não conseguiu abrir:\n\n{_e}\n\n"
                f"Detalhes salvos em: erro_fatal.txt (nesta mesma pasta)."
            )
            _raiz_erro.destroy()
        except Exception:
            pass
        sys.exit(1)

    if not janela.aprovado:
        sys.exit(0)


    _iniciar_tray(forcar_primeiro_ciclo=_esta_em_pausa())
